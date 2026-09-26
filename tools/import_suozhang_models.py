"""Gated local import of the audited 2026.9.13 N4.5 / N5 Word set.

Default: prepare candidates, replay the 8.31 baseline and write a review plan.
--apply: apply that frozen plan, checking all inputs and backing up every target.
No network, image generation, asset writes, or publication.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
import convert
import codex_update_match as match
import suozhang_r18_merge_match as merge

ROOT = Path(__file__).resolve().parents[1]
VERSION = '2026.9.13'
DEFAULT_OUT = ROOT / 'output' / 'suozhang-20260913-import'
SOURCE_HASHES = {
    'N4.5-regular': 'f080012c2b1b88e0644e2e79f07b927ce5c52a27bd9055990a5faddad1bf747d',
    'N4.5-upper': '60112b8143286e6b7b0b525466d7ef66903e31855f0784fcd522de1abd4e905f',
    'N4.5-lower': 'cc339b567c672b744eb76c8d4b6aa181d0591c4d99d99cfa561d2a5bb365a20d',
    'N5-regular': '7510a1267aab6cff5e39228376d8ad292e87c165351cc502209c9c4be818c82f',
    'N5-upper': 'adb708a5807cac4e5056c96021415328e7cf1e049af9306ad631f575042a6f76',
    'N5-lower': '488f3e254023419251d4cdd33781819bd811354c63d6e216955d87ca86ef59a6',
}
# Source hashes above make these reviewed paragraph decisions version-specific.
# Short prose without punctuation is indistinguishable from a title in isolation.
PROSE_PARAGRAPHS = {
    'N5-regular': (330, 334),
    'N5-upper': (334, 590, 591, 592, 775, 776, 777),
    'N5-lower': (),
}
BOOKS = ('suozhang', 'suozhang_r18', 'suozhang_nai5', 'suozhang_nai5_r18')
TITLES = {
    'suozhang': '所长N4.5常规NovelAI个人法典',
    'suozhang_r18': '所长N4.5色色NovelAI个人法典（合并版）',
    'suozhang_nai5': '所长N5常规NovelAI个人法典',
    'suozhang_nai5_r18': '所长N5色色NovelAI个人法典（合并版）',
}


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read_json(path):
    return json.loads(Path(path).read_text(encoding='utf-8'))


def write_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')


def data_hashes():
    return {str(p.relative_to(ROOT)):digest(p) for p in sorted((ROOT/'site/data').rglob('*.json'))}


def restore_food_emoji(doc, paragraph_index=6694):
    """Restore the six visually audited inline pictures, in their XML positions."""
    p = doc.paragraphs[paragraph_index]
    if len(convert.paragraph_blips(p)) != 6:
        raise ValueError('Food emoji paragraph no longer has six inline pictures')
    icons = dict(zip((f'/word/media/image{i}.png' for i in range(1,7)), '🍔🍟🌮🌭🍕🍚'))
    chunks = []
    for node in p._p.iter():
        if node.tag == qn('w:t'):
            chunks.append(node.text or '')
        elif node.tag == qn('a:blip'):
            part = doc.part.related_parts[node.get(qn('r:embed'))]
            chunks.append(icons[str(part.partname)])
    text = ''.join(chunks)
    if text.translate({ord(c):None for c in icons.values()}) != p.text:
        raise ValueError('Inline emoji restoration changed source text')
    return {paragraph_index: text}


def checked_normalize(entries, reference):
    result, audit = match.normalize_suozhang_entries(entries, reference)
    if audit['blockers']:
        raise ValueError('Role normalization blocked: '+json.dumps(audit['blockers'], ensure_ascii=True))
    return result, audit


def verify_recovered_prose(doc, entries, audit):
    """Each reclassified source line must survive in its positive/role fields."""
    def squeezed(value):
        return re.sub(r'\s+', '', value)
    bodies=[]
    for entry in entries:
        parts=[entry['tags']]+[p.get('prompt','') for p in entry.get('characterPrompts',[])]
        bodies.extend(squeezed(p) for p in parts)
    checked=0
    for record in audit:
        if record['after']!='tag':continue
        paragraph=doc.paragraphs[record['paragraph']]
        for line in convert.visible_lines(paragraph.text):
            if hashlib.sha256(line.encode('utf-8')).hexdigest()!=record['sha256']:continue
            body=re.sub(r'^(?:char\d*|角色\d+|人物\d+)\s*[:：]\s*','',line,flags=re.I)
            if body and not any(squeezed(body) in value for value in bodies):
                raise ValueError(f"Recovered prose was lost at paragraph {record['paragraph']}")
            checked+=1
    return checked


def compact_result(result):
    return {
        'summary': result['summary'],
        'validationIssues': result['validationIssues'],
        'reviewCount': len(result['review']),
        'changes': [{'oldId':r['old']['id'], 'newIndex':r['new']['index'], 'method':r['method'],
                     'fields':r['changes']} for r in result['matches'] if r['changes']],
        'additions':[{'index':e['index'], 'path':e['path']} for e in result['additions']],
        'removals':[e['id'] for e in result['removals']],
    }


def checked_match(old, new):
    result = match.match_entries(old, new)
    if result['validationIssues'] or result['review']:
        raise ValueError(f"Unresolved matching: {len(result['review'])} reviews; {result['validationIssues']}")
    return result


def validate_book(book):
    entries = book['entries']
    assert book['entryCount'] == len(entries)
    assert book['tree'] == convert.build_tree(entries)
    assert book['imagedCount'] == sum(bool(e.get('image')) for e in entries)
    assert len({e['id'] for e in entries}) == len(entries)
    assert all(e['title'] and e['title'] != '(未命名)' for e in entries)
    assert not any(match.STANDALONE_ROLE_TITLE_RE.fullmatch(e['title']) for e in entries)
    assert not any(re.search(r'(?im)^\s*char\d*\s*[:：]', e['tags']) for e in entries)
    assert all(e.get('tags') or e.get('characterPrompts') for e in entries)
    assert not any('法典相关网站' in e['path'] for e in entries)
    assert not any('负面提示词：' in e['tags'] for e in entries)
    normalized, audit = checked_normalize(entries, entries)
    assert all(match.norm_prompt_value(a) == match.norm_prompt_value(b) for a,b in zip(entries,normalized))
    assert len(normalized) == len(entries)
    assert audit['changedEntries'] == 0


def prepare(source, out):
    out.mkdir(parents=True, exist_ok=True)
    if (out/'applied.json').exists():
        raise ValueError('This frozen plan is already applied; use --verify')
    before = data_hashes()
    index = read_json(ROOT/'site/data/codexes.json')
    old_index = {e['id']:e for e in index}
    formal = {cid:read_json(ROOT/'site/data'/f'{cid}.json') for cid in BOOKS[:2]}
    if any(c['version']!='2026.8.31' for c in formal.values()):
        raise ValueError('This import expects the audited 8.31 formal baseline')
    if any(cid in old_index or (ROOT/'site/data'/f'{cid}.json').exists() for cid in BOOKS[2:]):
        raise ValueError('N5 book identity already exists; do not re-create its IDs')
    halves = merge.partition_formal_entries(formal['suozhang_r18']['entries'])
    baseline_paths = [ROOT/'output/所长常规-8.31-匹配测试/new-version-match.json',
                      ROOT/'output/所长色色-8.31-合并匹配测试/merged-source.json']
    base_regular,_ = match.load_baseline_json(baseline_paths[0])
    base_regular,_ = checked_normalize(base_regular, formal['suozhang']['entries'])
    base_halves,_ = merge.load_merged_source_snapshot(baseline_paths[1])
    base_r18=[]
    for half in ('upper','lower'):
        entries,_=checked_normalize(base_halves[half], halves[half])
        entries,_=merge.apply_known_source_title_corrections(entries, halves[half], half=half)
        base_r18.extend(entries)
    baseline={}
    for cid,entries in zip(BOOKS[:2],(base_regular,base_r18)):
        result=checked_match(formal[cid]['entries'],entries)
        if not result['summary']['strictReplayPass']:
            raise ValueError(f'{cid}: strict baseline replay failed')
        baseline[cid]=result['summary']

    parsed={}; source_records={}
    for path in sorted(source.glob('*.docx')):
        model='N5' if path.name.startswith('N5') else 'N4.5'
        half='regular' if '常规' in path.name else 'upper' if '（上）' in path.name else 'lower'
        key=f'{model}-{half}'
        if key in parsed or digest(path)!=SOURCE_HASHES.get(key):
            raise ValueError(f'Unreviewed or duplicate source: {path.name}')
        structure=match.inspect_docx_package(path)
        if not match.word_structure_safe(structure):
            raise ValueError(f'Unsupported Word structure: {key}')
        doc=Document(path); audit=[]
        overrides=restore_food_emoji(doc) if key=='N4.5-regular' else {}
        raw=convert.parse_standard_docx_items(doc, suozhang_model=model,
            kind_overrides={p:'tag' for p in PROSE_PARAGRAPHS.get(key,())},
            text_overrides=overrides, audit=audit)
        reference=(formal['suozhang']['entries'] if half=='regular' else halves[half]) if model=='N4.5' else []
        entries,roles=checked_normalize(raw,reference)
        recovered=verify_recovered_prose(doc,entries,audit) if model=='N5' else 0
        if model=='N4.5' and half!='regular':
            entries,_=merge.apply_known_source_title_corrections(entries,halves[half],half=half)
        unnamed=Counter()
        for entry in entries:
            if entry['title']=='(未命名)':
                category=entry['path'][-1]
                unnamed[category]+=1
                entry['title']=f"{category}（无标题{unnamed[category]}）"
        parsed[key]=entries
        source_records[key]={'path':str(path), 'sha256':digest(path), 'structure':structure,
            'parsed':len(entries), 'classificationChanges':audit,
            'recoveredProseLinesVerified':recovered,
            'artistCards':sum(merge.is_artist_group_entry(e) for e in entries),
            'ocCards':sum(merge.is_compiler_oc_entry(e) for e in entries),
            'roleBoxes':sum(len(e.get('characterPrompts',[])) for e in entries),
            'negativeEntries':sum(bool(e.get('negative')) for e in entries),
            'unnamedTitles':dict(unnamed), 'restoredInlineEmoji':6 if overrides else 0}
        print(json.dumps({'parsed':key,'count':len(entries),'roles':source_records[key]['roleBoxes']},ensure_ascii=True),flush=True)
    if set(parsed)!=set(SOURCE_HASHES):
        raise ValueError('Expected all six source documents')

    merges={}
    for model in ('N4.5','N5'):
        upper=parsed[f'{model}-upper']; lower=parsed[f'{model}-lower']
        duplicate_oc=[]
        if model=='N5':
            upper_oc=Counter(merge.fingerprint(e) for e in upper if merge.is_compiler_oc_entry(e))
            kept=[]
            for entry in lower:
                sig=merge.fingerprint(entry)
                if merge.is_compiler_oc_entry(entry) and upper_oc[sig]>0:
                    upper_oc[sig]-=1; duplicate_oc.append(entry['title'])
                else:
                    kept.append(entry)
            lower=kept
        merged=merge.merge_source_halves(upper,lower)
        parsed[f'{model}-r18']=merged['merged']
        merges[model]={**merged['stats'],**merged['special'],'duplicateLowerOcRemoved':len(duplicate_oc)}
        write_json(out/f'{model}-merged-source.json',{'codexId':'suozhang_r18' if model=='N4.5' else 'suozhang_nai5_r18',
            'version':VERSION,'entryCount':len(merged['merged']),'entries':merged['merged'],
            'stats':merged['stats'],'specialHandling':merged['special']})

    generated={}; results={}
    for cid,key in zip(BOOKS,('N4.5-regular','N4.5-r18','N5-regular','N5-r18')):
        source_entries=parsed[key]
        old=formal.get(cid) or {'id':cid,'title':TITLES[cid],'author':'戒红所','version':VERSION,'entries':[]}
        result=checked_match(old['entries'],source_entries)
        if cid in BOOKS[2:]:
            # First publication of a distinct book: every entry belongs to its
            # first update batch; no old-model IDs or image lookup is involved.
            for entry in source_entries:
                entry['isNew']=True
        builder=merge.build_applied_codex if cid=='suozhang_r18' else match.build_applied_codex
        reserved=merge.collect_reserved_asset_ids() if cid=='suozhang_r18' else {
            p.stem for base in (ROOT/'site/images'/cid,ROOT/'originals'/cid) if base.is_dir() for p in base.iterdir() if p.is_file()}
        applied,stats=builder(old,source_entries,result,VERSION,reserved_ids=reserved,
                             previous_update_filters=old_index.get(cid,{}).get('updateFilters'))
        applied['title']=TITLES[cid]
        # First-time N5 r18 uses one fresh book namespace, not the N4.5 halves.
        for e in applied['entries']:
            e.pop('sourceHalf',None); e.pop('sourceIndex',None)
            if cid in BOOKS[2:]:
                assert e.get('image') is None and not e.get('original')
        validate_book(applied)
        by_id={e['id']:e for e in applied['entries']}
        for old_entry in old['entries']:
            if old_entry['id'] in by_id:
                new_entry=by_id[old_entry['id']]
                for k,v in old_entry.items():
                    if k not in {*match.SOURCE_CONTENT_KEYS,'updateBatches','sourceHalf','sourceIndex'}:
                        assert new_entry.get(k)==v,(cid,old_entry['id'],k)
        generated[cid]=applied
        results[cid]={'match':compact_result(result),'apply':stats}
        write_json(out/'candidates'/f'{cid}.json',applied)
        if cid in old_index:
            index=match.build_updated_codex_index(index,applied,cid)
            next(e for e in index if e['id']==cid)['title']=TITLES[cid]
        else:
            meta={'id':cid,'type':'codex','title':TITLES[cid],'version':VERSION,'author':'戒红所',
                  'entryCount':applied['entryCount'],'imagedCount':0,'hasOriginal':False,
                  'source':'戒红所 · N5所长法典','contributors':[{'name':'戒红所','role':'法典原作者'}],
                  'newFilterLabel':match.update_filter_label(VERSION),
                  'updateFilters':match.update_filter_history([],VERSION)}
            if cid.endswith('_r18'):meta['nsfw']=True
            adjacent='suozhang_r18' if cid.endswith('_r18') else 'suozhang'
            index.insert(next(i for i,e in enumerate(index) if e['id']==adjacent)+1,meta)
        match.validate_update_batch_contract(applied,next(e for e in index if e['id']==cid)['updateFilters'])
    assert all(item==next(e for e in index if e['id']==item['id']) for item in old_index.values() if item['id'] not in BOOKS)
    write_json(out/'candidates/codexes.json',index)
    write_json(out/'regular-source.json',{'version':VERSION,'entries':parsed['N4.5-regular']})
    write_json(out/'N5-regular-source.json',{'version':VERSION,'entries':parsed['N5-regular']})
    payloads={f'site/data/{cid}.json':out/'candidates'/f'{cid}.json' for cid in (*BOOKS,'codexes')}
    if data_hashes()!=before:
        raise ValueError('Formal data changed during planning; re-run against the latest data')
    plan={'version':VERSION,'baseline':baseline,'sources':source_records,'merges':merges,'books':results,
          'beforeDataHashes':before,
          'targets':{rel:{'candidate':str(path),'sha256':digest(path),'beforeSha256':before.get(str(Path(rel)))} for rel,path in payloads.items()}}
    write_json(out/'plan.json',plan)
    print(json.dumps({cid:v['apply']|{'newIds':'(see plan)'} for cid,v in results.items()},ensure_ascii=True,indent=2))


def apply_plan(out):
    plan=read_json(out/'plan.json')
    if (out/'applied.json').exists():
        raise ValueError('Already applied; use --verify')
    if data_hashes()!=plan['beforeDataHashes']:
        raise ValueError('Formal data changed after plan creation')
    for source in plan['sources'].values():
        if digest(source['path'])!=source['sha256']:
            raise ValueError('Source changed after plan creation')
    pending=[]
    for rel,item in plan['targets'].items():
        candidate=Path(item['candidate']); target=ROOT/rel
        if digest(candidate)!=item['sha256']:
            raise ValueError(f'Candidate changed: {rel}')
        original=target.read_bytes() if target.exists() else None
        if original is not None:
            backup=out/'before'/target.name
            backup.parent.mkdir(parents=True,exist_ok=True)
            if backup.exists() and backup.read_bytes()!=original:
                raise ValueError('Refusing to overwrite a different backup')
            backup.write_bytes(original)
        pending.append((target,candidate.read_bytes(),original))
    completed=[]
    try:
        for target,content,original in pending:
            tmp=target.with_suffix('.json.tmp')
            tmp.write_bytes(content); tmp.replace(target)
            completed.append((target,original))
    except BaseException:
        for target,original in reversed(completed):
            if original is None:
                target.unlink(missing_ok=True)
            else:
                restore=target.with_suffix('.json.rollback')
                restore.write_bytes(original); restore.replace(target)
        raise
    write_json(out/'applied.json',{'version':plan.get('version',VERSION),'files':{str(p.relative_to(ROOT)):digest(p) for p,_,_ in pending}})
    verify(out)


def verify(out):
    plan=read_json(out/'plan.json')
    applied=read_json(out/'applied.json')
    for rel,target in plan['targets'].items():
        assert digest(ROOT/rel)==target['sha256']==digest(target['candidate']),rel
        assert applied['files'][str(Path(rel))]==target['sha256'],rel
    index=read_json(ROOT/'site/data/codexes.json')
    for cid in BOOKS:
        book=read_json(ROOT/'site/data'/f'{cid}.json')
        validate_book(book)
        assert book==read_json(out/'candidates'/f'{cid}.json')
        meta=next(e for e in index if e['id']==cid)
        for key in ('title','version','entryCount','imagedCount'):
            assert book[key]==meta[key]
        match.validate_update_batch_contract(book,meta['updateFilters'])
    assert index==read_json(out/'candidates/codexes.json')
    changed={str(Path(k)) for k in plan['targets']}
    allowed_derived={'site/data/updates.json','site/data/share-index.json'}
    unchanged=0
    for rel,sha in plan['beforeDataHashes'].items():
        posix=Path(rel).as_posix()
        if rel in changed or posix.startswith(('site/data/share/','site/data/tag_zh/')) or posix in allowed_derived:
            continue
        assert digest(ROOT/rel)==sha,rel
        unchanged+=1
    print(json.dumps({'verifiedBooks':list(BOOKS),'unrelatedDataFilesUnchanged':unchanged},ensure_ascii=True))


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source-dir',type=Path,default=ROOT.parent/'新数据'/'（解压密码：1984）所长NovelAI个人法典（2026.9.13版，一般所长整理）')
    parser.add_argument('--out-dir',type=Path,default=DEFAULT_OUT)
    mode=parser.add_mutually_exclusive_group()
    mode.add_argument('--apply',action='store_true')
    mode.add_argument('--verify',action='store_true')
    args=parser.parse_args()
    if args.apply:apply_plan(args.out_dir)
    elif args.verify:verify(args.out_dir)
    else:prepare(args.source_dir,args.out_dir)


if __name__=='__main__':
    main()
