import sys
import unittest
from pathlib import Path
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parent))
import convert
import codex_update_match as match
import suozhang_r18_merge_match as merge


def document(category='场景'):
    doc=Document()
    doc.styles.add_style('toc 1',WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph('目录',style='toc 1')
    heading=doc.add_paragraph(category)
    outline=OxmlElement('w:outlineLvl'); outline.set(qn('w:val'),'0')
    heading._p.get_or_add_pPr().append(outline)
    return doc


class SuozhangModelTests(unittest.TestCase):
    def test_reference_websites_are_not_prompt_cards(self):
        doc=document('法典相关网站')
        doc.add_paragraph('教程网站')
        doc.add_paragraph('https://example.com/tutorial')
        heading=doc.add_paragraph('场景')
        outline=OxmlElement('w:outlineLvl'); outline.set(qn('w:val'),'0')
        heading._p.get_or_add_pPr().append(outline)
        doc.add_paragraph('林间散步')
        doc.add_paragraph('girl, walking,')
        for model in ('N4.5','N5'):
            entries=convert.parse_standard_docx_items(doc,suozhang_model=model)
            self.assertEqual(len(entries),1)
            self.assertEqual(entries[0]['title'],'林间散步')
            self.assertEqual(entries[0]['path'],['场景'])

    def test_chinese_and_english_prose_and_short_roles_survive(self):
        doc=document()
        for text in ('林间散步','场景：树林,石阶','char1：girl','微风吹动头发,手里拿着一本书',
                     'char2：boy','expression: a gentle smile.'):
            doc.add_paragraph(text)
        raw=convert.parse_standard_docx_items(doc,suozhang_model='N5')
        entries,audit=match.normalize_suozhang_entries(raw,[])
        self.assertEqual(len(entries),1)
        self.assertFalse(audit['blockers'])
        self.assertEqual(entries[0]['title'],'林间散步')
        self.assertEqual(entries[0]['tags'],'场景：树林,石阶')
        self.assertEqual(entries[0]['characterPrompts'],[
            {'label':'char1','prompt':'girl\n微风吹动头发,手里拿着一本书'},
            {'label':'char2','prompt':'boy\nexpression: a gentle smile.'}])

    def test_panel_labels_remain_in_one_card(self):
        doc=document()
        lines=['四格散步','An illustration of a peaceful afternoon in a green park.',
               'top-left panel:','no makeup.','top-right panel:','A person walks past a tree.']
        for text in lines:doc.add_paragraph(text)
        entries=convert.parse_standard_docx_items(doc,suozhang_model='N5')
        self.assertEqual(len(entries),1)
        self.assertEqual(entries[0]['tags'],'\n'.join(lines[1:]))

    def test_parenthetical_title_note_is_not_prompt(self):
        doc=document()
        doc.add_paragraph('水彩画(不加质量词,保留纸张质感)')
        doc.add_paragraph('watercolor, paper texture,')
        entries=convert.parse_standard_docx_items(doc,suozhang_model='N5')
        self.assertEqual(entries[0]['title'],'水彩画(不加质量词,保留纸张质感)')

    def test_artist_continuations_and_negative_stay_together(self):
        doc=document('编纂者常用画师组')
        for text in ('NAI5时期','1，artist:sample,','-1::flat color::,','year 2026,',
                     '负面提示词：watermark,logo,','NAI4.5时期：','1，artist:old,','2，artist:other,'):
            doc.add_paragraph(text)
        for model in ('N4.5','N5'):
            entries=convert.parse_standard_docx_items(doc,suozhang_model=model)
            self.assertEqual(len(entries),3)
            self.assertEqual(entries[0]['tags'],'artist:sample,\n-1::flat color::,\nyear 2026,')
            self.assertEqual(entries[0]['negative'],'watermark,logo,')
            self.assertNotIn('negative',entries[1])
            self.assertEqual(entries[2]['title'],'NAI4.5时期：2')

    def test_negative_is_part_of_match_and_replay(self):
        old={'id':'demo-0001','path':['画风'],'title':'示例','tags':'artist:sample,','negative':'logo,','image':'old.jpg'}
        new={**old,'negative':'watermark,'}
        result=match.match_entries([old],[new])
        self.assertEqual(result['summary']['contentChanged'],1)
        self.assertFalse(result['summary']['strictReplayPass'])
        applied,_=match.build_applied_codex({'id':'demo','version':'2026.8.31','entries':[old]},[new],result,'2026.9.13')
        self.assertEqual(applied['entries'][0]['negative'],'watermark,')
        self.assertEqual(applied['entries'][0]['image'],'old.jpg')
        self.assertEqual(result['matches'][0]['new']['negative'],'watermark,')

    def test_different_negative_must_not_be_deduplicated(self):
        entry={'path':['编纂者常用画师组'],'title':'NAI5时期','tags':'artist:sample,','negative':'logo,'}
        with self.assertRaises(ValueError):
            merge.merge_source_halves([entry],[{**entry,'negative':'watermark,'}])

    def test_n5_cannot_resolve_to_legacy_identity(self):
        self.assertEqual(convert.codex_id('N5所长常规NovelAI个人法典'),'suozhang_nai5')
        with self.assertRaises(ValueError):
            convert.codex_id('N5所长色色NovelAI个人法典（上）')

    def test_version_scoped_short_prose_override(self):
        doc=document()
        doc.add_paragraph('林间散步'); doc.add_paragraph('girl, walking,')
        doc.add_paragraph('风吹树叶沙沙作响')
        entries=convert.parse_standard_docx_items(doc,suozhang_model='N5',kind_overrides={4:'tag'})
        self.assertEqual(entries[0]['tags'],'girl, walking,\n风吹树叶沙沙作响')


if __name__=='__main__':
    unittest.main()
