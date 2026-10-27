"""所长法典版本更新的通用入口（2026.9.25 起替代逐版专用脚本）。

  survey    只读调研新版 Word，写出调研报告和版本配置草稿
  preview   按版本配置算出结果给人核对（只读，不冻结）
  prepare   按 tools/suozhang_versions/<版本>.json 生成冻结候选与计划（只写 output/）
  apply     校验冻结计划、备份后写四本法典 JSON 与书目
  verify    只读复验已应用的计划
  selfcheck 用 9.25 的存档输入重放，确认改过代码后结果不变
  compare   两份配置在同一份底数据上重建、逐条比对（演练对标准答案；平时对改动前的配置）

不写图片、不联网、不发布。完整流程见本地私有文档 docs/经验/所长法典更新手册.md。
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
from collections import Counter, defaultdict
from pathlib import Path

from docx import Document
import convert
import codex_update_match as match
import suozhang_r18_merge_match as merge
import import_suozhang_models as base

ROOT = base.ROOT
VERSIONS_DIR = ROOT / 'tools' / 'suozhang_versions'
NEW_DATA = ROOT.parent / '新数据'
BOOKS = base.BOOKS
REGULAR_BOOKS = ('suozhang', 'suozhang_nai5')
BOOK_OF = {('N4.5', False): 'suozhang', ('N4.5', True): 'suozhang_r18',
           ('N5', False): 'suozhang_nai5', ('N5', True): 'suozhang_nai5_r18'}
SNAPSHOTS = {'suozhang': 'regular-source.json', 'suozhang_r18': 'N4.5-merged-source.json',
             'suozhang_nai5': 'N5-regular-source.json', 'suozhang_nai5_r18': 'N5-merged-source.json'}
HALF_PREFIXES = merge.HALF_PREFIXES
PART_RE = re.compile(r'^N(?:4\.5|5)-(?:regular|upper|lower|v\d+)$')
VERSION_RE = re.compile(r'(\d{4})\.(\d{1,2})\.(\d{1,2})版')
# Small cards that belong to the card above them (画风搜集 since 9.25).
# 「正面」不在内：它也是常见的视角变体标题（正面 / 背面），9.25 的群友收录里就有。
LABEL_TITLES = frozenset({'负面', '负面提示词', '设置', '参数', '附带tag', '附带'})
# Short titles that are real variant cards, not prose continuations.
VARIANT_TITLE_RE = re.compile(r'^(?:原版|其他版本|其它版本|版本|差分|变体)')
FRESH_PROMPT_RE = re.compile(r'^[\[{(0-9A-Za-z]')
CHAR_LINE_RE = re.compile(r'^char\d*\s*[:：]', re.I)
ON_IMAGE_TEXT_RE = re.compile(r'\btext\s*[:：]', re.I)
UNTITLED_RE = re.compile(r'（无标题\d+）$')
CN_DIGITS = dict(zip('一二三四五六七八九', range(1, 10)))


# ---------------------------------------------------------------- basics

def sha256_text(text):
    return hashlib.sha256(text.encode('utf-8')).hexdigest()


def path_sha(path):
    return sha256_text(json.dumps(list(path), ensure_ascii=False))


def version_tuple(version):
    return tuple(int(part) for part in str(version).split('.'))


def version_slug(version):
    year, month, day = version_tuple(version)
    return f'{year}{month:02d}{day:02d}'


def default_out(version):
    return ROOT / 'output' / f'suozhang-{version_slug(version)}-import'


def cn_number(text):
    if text == '十':
        return 10
    if '十' in text:
        tens, _, ones = text.partition('十')
        return CN_DIGITS.get(tens, 1) * 10 + (CN_DIGITS[ones] if ones else 0)
    return CN_DIGITS[text]


def part_key(name):
    upper = name.upper()
    if upper.startswith('N4.5'):
        model = 'N4.5'
    elif upper.startswith('N5'):
        model = 'N5'
    else:
        raise ValueError(f'无法识别模型（N4.5 / N5）：{name}')
    if '常规' in name:
        return f'{model}-regular'
    volume = re.search(r'[（(]卷([一二三四五六七八九十]+)[）)]', name)
    if volume:
        return f'{model}-v{cn_number(volume.group(1))}'
    for marker, half in (('上', 'upper'), ('下', 'lower')):
        if re.search(rf'[（(]{marker}[）)]', name):
            return f'{model}-{half}'
    raise ValueError(f'无法识别分卷：{name}')


def part_model(key):
    return key.split('-', 1)[0]


def part_book(key):
    model, kind = key.split('-', 1)
    return BOOK_OF[(model, kind != 'regular')]


def part_rank(key):
    kind = key.split('-', 1)[1]
    ranks = {'regular': 0, 'upper': 1, 'lower': 2}
    return ranks[kind] if kind in ranks else int(kind[1:])


def ordered_parts(parts, book):
    keys = sorted((k for k in parts if part_book(k) == book), key=part_rank)
    kinds = {k.split('-', 1)[1][0] for k in keys}
    if 'v' in kinds and kinds & {'u', 'l'}:
        raise ValueError(f'{book}: 分卷命名混用了「卷N」和「上/下」')
    return keys


def source_version(names):
    found = {'.'.join(str(int(x)) for x in m.groups()) for name in names for m in [VERSION_RE.search(name)] if m}
    if len(found) != 1:
        raise ValueError(f'无法从文件名确定唯一版本号：{sorted(found)}')
    return found.pop()


def special(entry):
    if merge.is_artist_group_entry(entry):
        return 'artist'
    if merge.is_compiler_oc_entry(entry):
        return 'oc'
    return None


def body(entry):
    return match.norm_prompt_value(entry)


def alias_path(path, pairs):
    for old, new in sorted(pairs, key=lambda pair: -len(pair[0])):
        if tuple(path[:len(old)]) == tuple(old):
            return list(new) + list(path[len(old):])
    return list(path)


def aliased(entries, pairs):
    return [dict(entry, path=alias_path(entry.get('path', []), pairs)) for entry in entries]


def reserved_ids(book, cap=None):
    """Asset stems already used by a book; ``cap`` keeps only IDs that existed then."""
    if book == 'suozhang_r18':
        ids = merge.collect_reserved_asset_ids()
    else:
        ids = {p.stem for folder in (ROOT / 'site/images' / book, ROOT / 'originals' / book)
               if folder.is_dir() for p in folder.iterdir() if p.is_file()}
    if cap is None:
        return ids
    kept = set()
    for value in ids:
        prefix, _, number = value.rpartition('-')
        if not number.isdigit() or int(number) <= cap.get(prefix, 0):
            kept.add(value)
    return kept


def id_caps(entries):
    caps = {}
    for entry in entries:
        prefix, _, number = str(entry.get('id', '')).rpartition('-')
        if number.isdigit():
            caps[prefix] = max(caps.get(prefix, 0), int(number))
    return caps


# ---------------------------------------------------------------- config

def load_config(path):
    config = base.read_json(path)
    for key in ('version', 'outDir', 'parts'):
        if key not in config:
            raise ValueError(f'{path}: 缺少字段 {key}')
    for key, part in config['parts'].items():
        if not PART_RE.fullmatch(key):
            raise ValueError(f'{path}: 分卷键不合法 {key}')
        half = part.get('half')
        if part_book(key) in REGULAR_BOOKS:
            if half is not None:
                raise ValueError(f'{path}: 常规本 {key} 不能有 half')
        elif half not in merge.HALF_PREFIXES:
            raise ValueError(f'{path}: {key} 的 half 必须是 upper 或 lower')
    if not config.get('historical'):
        for key in ('baseVersion', 'source', 'baseline', 'expected'):
            if key not in config:
                raise ValueError(f'{path}: 缺少字段 {key}')
    config['_path'] = str(path)
    return config


def config_path(version):
    return VERSIONS_DIR / f'{version}.json'


def previous_config(version):
    older = []
    for path in VERSIONS_DIR.glob('*.json'):
        candidate = path.stem
        if re.fullmatch(r'\d{4}\.\d{1,2}\.\d{1,2}', candidate) and version_tuple(candidate) < version_tuple(version):
            older.append(candidate)
    if not older:
        raise ValueError('tools/suozhang_versions/ 里没有更早版本的配置')
    return load_config(config_path(max(older, key=version_tuple)))


def aliases_of(config, book):
    return [(tuple(old), tuple(new)) for old, new in config.get('pathAliases', {}).get(book, [])]


def install_overrides(config):
    version = config['version']
    for book, table in config.get('newOverrides', {}).items():
        key = (book, version)
        current = match.AUDITED_NEW_OVERRIDES.get(key)
        if current is not None and current != table:
            raise ValueError(f'{book} {version}: 代码里已有不同的 NEW 修正表')
        match.AUDITED_NEW_OVERRIDES[key] = table


# ---------------------------------------------------------------- parsing

def source_files(source_dir):
    files = {}
    for path in sorted(Path(source_dir).glob('*.docx')):
        if path.name.startswith('~$'):  # Word lock file while the document is open
            continue
        key = part_key(path.name)
        if key in files:
            raise ValueError(f'同一分卷出现两份文件：{key}')
        files[key] = path
    if not files:
        raise ValueError(f'目录里没有 docx：{source_dir}')
    return files


TEMPLATE_LABELS = {'原tag': 'original', '翻译': 'translation', 'tag版本改编后': 'adapted', '改编后': 'adapted'}
TEMPLATE_LABEL_RE = re.compile(r'^(原Tag|翻译|Tag版本改编后|改编后)\s*[:：]\s*(.*)$', re.I | re.S)


def chapter_bounds(doc, name):
    paragraphs = doc.paragraphs
    starts = [i for i, p in enumerate(paragraphs) if p.text.strip() == name and convert.outline_lvl(p) == 0]
    if len(starts) != 1:
        raise ValueError(f'登记的特别章节「{name}」找到 {len(starts)} 处（应为 1 处）')
    start = starts[0]
    end = next((i for i in range(start + 1, len(paragraphs))
                if convert.outline_lvl(paragraphs[i]) == 0 and paragraphs[i].text.strip()), len(paragraphs))
    later = [p.text.strip() for p in paragraphs[end:] if convert.outline_lvl(p) == 0 and p.text.strip()]
    return start, end, later


def parse_template_chapter(doc, rule):
    """Read a chapter the parser skips: repeated 原Tag / 翻译 (/ 改编后) blocks, no titles.

    Each filled block becomes「<前缀>（N）」with the translation as its note; an
    adapted version becomes「<前缀>（N）：改编后」. Empty template blocks are dropped.
    NEW follows the author's highlight, as everywhere else.
    """
    start, end, later = chapter_bounds(doc, rule['chapter'])
    blocks, current, field = [], None, None
    for index in range(start + 1, end):
        paragraph = doc.paragraphs[index]
        if convert.outline_lvl(paragraph) is not None and paragraph.text.strip():
            raise ValueError(f"特别章节「{rule['chapter']}」里出现了小节标题，模板变了，停下复核")
        for line in convert.visible_lines(paragraph.text):
            label = TEMPLATE_LABEL_RE.match(line)
            if label:
                field = TEMPLATE_LABELS[label.group(1).lower()]
                if field == 'original':
                    current = {'original': [], 'translation': [], 'adapted': [], 'pink': False}
                    blocks.append(current)
                elif current is None:
                    raise ValueError(f"特别章节「{rule['chapter']}」第 {index} 段在第一个「原Tag」之前")
                rest = label.group(2).strip()
                if rest:
                    current[field].append(rest)
            elif current is None:
                raise ValueError(f"特别章节「{rule['chapter']}」第 {index} 段在第一个「原Tag」之前")
            else:
                current[field].append(line)
            current['pink'] = current['pink'] or convert.is_pink(paragraph)
    items, number, empty = [], 0, 0
    for block in blocks:
        if not block['original']:
            if block['translation'] or block['adapted']:
                raise ValueError(f"特别章节「{rule['chapter']}」有只填了翻译 / 改编、没有原文的格，停下复核")
            empty += 1
            continue
        number += 1
        title = f"{rule['titlePrefix']}（{number}）"
        item = {'title': title, 'path': [rule['chapter']], 'tags': '\n'.join(block['original']),
                'isNew': block['pink']}
        if block['translation']:
            item['note'] = '\n'.join(block['translation'])
        items.append(item)
        if block['adapted']:
            items.append({'title': f'{title}：改编后', 'path': [rule['chapter']],
                          'tags': '\n'.join(block['adapted']), 'isNew': block['pink']})
    stats = {'blocks': len(blocks), 'filled': number, 'emptyTemplates': empty,
             'adapted': sum(bool(b['adapted']) for b in blocks), 'items': len(items)}
    return items, later, stats


def insert_chapter_items(raw, items, later_sections):
    later = set(later_sections)
    position = next((i for i, entry in enumerate(raw) if entry.get('path', [])[:1] and entry['path'][0] in later),
                    len(raw))
    return raw[:position] + items + raw[position:]


def translation_card_starts(doc, key, kinds):
    """Prompt lines for each registered translation block (``kinds`` value 'translation').

    Consecutive translation paragraphs are one block; the prompt is the run
    of prompt lines right above it, stopping at a blank line, a heading or
    another block.  The prompt starts an untitled card of its own unless a
    card title sits right above it.  Returns (prompt start paragraphs that
    begin a card, first paragraphs of blocks with no prompt above them).
    """
    classify = convert.classify_nai5_line if part_model(key) == 'N5' else convert.classify
    paragraphs = doc.paragraphs

    def kind(index):
        return kinds.get(index) or classify(paragraphs[index].text.strip())

    def boundary(index):
        level = convert.outline_lvl(paragraphs[index])
        return not paragraphs[index].text.strip() or (level is not None and level <= 3) or kinds.get(index) == 'translation'

    starts, misfits = [], []
    for index in sorted(i for i, value in kinds.items() if value == 'translation'):
        if kinds.get(index - 1) == 'translation':
            continue
        top = index
        while top > 0 and not boundary(top - 1) and kind(top - 1) == 'tag':
            top -= 1
        if top == index:
            misfits.append(index)
        elif top == 0 or boundary(top - 1) or kind(top - 1) != 'title':
            starts.append(top)
    return starts, misfits


def prune_translations(docs, prose, titles, translations):
    """Drop translation blocks that no longer sit under a prompt; returns what was dropped per part."""
    dropped = {}
    for key, registered in list(translations.items()):
        kinds = {**{int(i): 'tag' for i in prose.get(key, {})}, **{int(i): 'title' for i in titles.get(key, {})},
                 **{int(i): 'translation' for i in registered}}
        _, misfits = translation_card_starts(docs[key], key, kinds)
        for start in misfits:
            index = start
            while str(index) in registered:
                registered.pop(str(index))
                dropped.setdefault(key, []).append(index)
                index += 1
        if not registered:
            translations.pop(key)
    return dropped


def load_part(path, key, prose=None, emoji_index=None, chapters=(), titles=None, translations=None, joined=None):
    """``prose`` paragraphs are read as prompt text, ``titles`` as card titles.

    A ``joined`` paragraph has its card title written straight into the
    prompt line; it is cut ``at`` that many characters (a ：/: right after
    the cut is dropped) into a title line and a prompt line.

    ``translations`` paragraphs are the Chinese translation of the prompt
    lines right above them and go into that card's note (see
    ``translation_card_starts``).
    """
    doc = Document(path)
    prose = {int(index): text_sha for index, text_sha in (prose or {}).items()}
    titles = {int(index): text_sha for index, text_sha in (titles or {}).items()}
    translations = {int(index): text_sha for index, text_sha in (translations or {}).items()}
    kinds = {}
    for registered, kind in ((prose, 'tag'), (titles, 'title'), (translations, 'translation')):
        for index in registered:
            if index in kinds:
                raise ValueError(f'{key}: 第 {index} 段登记了两种读法')
            kinds[index] = kind
    starts, misfits = translation_card_starts(doc, key, kinds)
    if misfits:
        raise ValueError(f'{key}: 第 {misfits} 段登记成翻译，但它上面不是提示词')
    kinds.update((index, 'card') for index in starts)
    for registered, name in ((prose, '续行'), (titles, '标题'), (translations, '翻译')):
        for index, text_sha in registered.items():
            if sha256_text(doc.paragraphs[index].text) != text_sha:
                raise ValueError(f'{key}: 第 {index} 段不是已审的{name}原文，段落位置变了')
    overrides = base.restore_food_emoji(doc, emoji_index) if emoji_index is not None else {}
    classify = convert.classify_nai5_line if part_model(key) == 'N5' else convert.classify
    for index, row in {int(i): row for i, row in (joined or {}).items()}.items():
        text = doc.paragraphs[index].text
        if index in kinds or index in overrides or sha256_text(text) != row['sha256']:
            raise ValueError(f'{key}: 第 {index} 段不是已审的「标题连着提示词」原文，或另有登记')
        title, rest = text[:row['at']].strip(), text[row['at']:].lstrip('：: ')
        first = rest.splitlines()[0].strip() if rest.strip() else ''
        inside_word = re.match(r'[A-Za-z0-9]{2}', text[row['at'] - 1:row['at'] + 1] or '')
        if (not title or inside_word or re.search(r'[,，:：{}\[\]\n]', title) or classify(title) != 'title'
                or not first or first[0] in ',，' or classify(first) != 'tag'):
            raise ValueError(f'{key}: 第 {index} 段在第 {row["at"]} 个字处断开，得不到「标题 + 提示词」')
        overrides[index] = title + '\n' + rest
    audit = []
    raw = convert.parse_standard_docx_items(
        doc, suozhang_model=part_model(key), kind_overrides=kinds, text_overrides=overrides, audit=audit)
    chapter_stats = {}
    for rule in chapters:
        if rule['part'] == key:
            items, later, chapter_stats[rule['chapter']] = parse_template_chapter(doc, rule)
            raw = insert_chapter_items(raw, items, later)
    return {'doc': doc, 'raw': raw, 'audit': audit, 'prose': sorted(prose), 'titles': sorted(titles),
            'translations': sorted(translations), 'joined': sorted(int(i) for i in (joined or {})),
            'chapters': chapter_stats,
            'emoji': emoji_index, 'sha256': base.digest(path), 'path': str(path)}


def fold_style_collection(entries, rule):
    """Fold 负面/设置/附带tag cards into the 画风 card above them.

    负面 → negative; 设置 → note (shown as 备注, never copied); 附带tag → the
    sample-scene prompt appended to the card, its role lines as role boxes.
    """
    path = list(rule['path'])
    out, seen = [], {}
    for entry in entries:
        if entry.get('path') != path:
            out.append(dict(entry))
            continue
        if entry['title'] not in LABEL_TITLES:
            if entry['title'] in rule['hosts']:
                if entry['title'] in seen:
                    raise ValueError(f"画风栏目里有两张同名卡：{entry['title']}")
                seen[entry['title']] = []
            out.append(dict(entry))
            continue
        host = out[-1] if out else None
        if not host or host.get('path') != path or host['title'] not in seen:
            raise ValueError(f"小标签前面没有所属画风卡：{entry['title']}")
        label = entry['title']
        if label in ('负面', '负面提示词'):
            if host.get('negative') or entry.get('characterPrompts'):
                raise ValueError('负面小标签的形状变了')
            host['negative'] = entry['tags']
        elif label in ('设置', '参数'):
            if entry.get('characterPrompts'):
                raise ValueError('设置小标签的形状变了')
            host['note'] = ' · '.join(re.sub(r'\s+', ' ', line).strip()
                                      for line in entry['tags'].splitlines() if line.strip())
        else:
            if entry['tags']:
                host['tags'] = host['tags'].rstrip() + '\n' + entry['tags']
            if entry.get('characterPrompts'):
                host['characterPrompts'] = list(host.get('characterPrompts') or []) + list(entry['characterPrompts'])
        host['isNew'] = bool(host.get('isNew') or entry.get('isNew'))
        seen[host['title']].append(label)
    expected = {title: list(labels) for title, labels in rule['hosts'].items()}
    if seen != expected:
        raise ValueError(f'画风栏目形状变了：{seen}')
    return out


def name_untitled(entries):
    counts = Counter()
    for entry in entries:
        if entry['title'] == '(未命名)':
            category = entry['path'][-1]
            counts[category] += 1
            entry['title'] = f'{category}（无标题{counts[category]}）'
    return dict(counts)


def source_split_hit(entries, rule):
    """Locate a registered untitled paragraph that the parser glued onto the card above."""
    hits = [(i, n) for i, entry in enumerate(entries)
            for n, line in enumerate(str(entry.get('tags', '')).split('\n'))
            if sha256_text(line.strip()) == rule['lineSha256']]
    if len(hits) != 1:
        return None
    index, line = hits[0]
    host = entries[index]
    if line == 0 or host.get('characterPrompts') or host.get('negative'):
        return None
    return index, line


def apply_source_splits(key, entries, rules):
    """Registered source errors, one exact paragraph each; nothing else is touched.

    The parser keeps an untitled paragraph on the card above it.  Instead of
    teaching the parser a general rule, each known case is listed in the
    version config by the hash of that paragraph and split back here.
    """
    for rule in rules:
        if rule['part'] != key:
            continue
        hit = source_split_hit(entries, rule)
        if hit is None:
            raise ValueError(f"已登记的源稿问题对不上了（{rule.get('note', '')}）：作者可能已修复或改了这段，"
                             '核实后从配置的 sourceSplits 删掉这条')
        index, line = hit
        host = entries[index]
        lines = host['tags'].split('\n')
        entries[index] = dict(host, tags='\n'.join(lines[:line]))
        entries.insert(index + 1, {'title': rule['title'], 'path': list(host['path']),
                                   'tags': '\n'.join(lines[line:]), 'isNew': bool(rule['isNew'])})
    return entries


def normalize_part(key, loaded, reference, style_rules=(), splits=()):
    entries, roles = base.checked_normalize(loaded['raw'], reference)
    recovered = base.verify_recovered_prose(loaded['doc'], entries, loaded['audit'])
    for rule in style_rules:
        if rule['part'] == key:
            entries = fold_style_collection(entries, rule)
    entries = apply_source_splits(key, entries, splits)
    untitled = name_untitled(entries)
    record = {
        'path': loaded['path'], 'sha256': loaded['sha256'], 'parsed': len(entries),
        'proseParagraphs': loaded['prose'], 'titleParagraphs': loaded['titles'],
        'translationParagraphs': loaded['translations'], 'joinedTitles': loaded['joined'],
        'foodEmojiParagraph': loaded['emoji'],
        'recoveredProseLinesVerified': recovered,
        'artistCards': sum(special(e) == 'artist' for e in entries),
        'ocCards': sum(special(e) == 'oc' for e in entries),
        'roleBoxes': sum(len(e.get('characterPrompts', [])) for e in entries),
        'standaloneRoleCardMerges': len(roles['standaloneRoleCardMerges']),
        'negativeEntries': sum(bool(e.get('negative')) for e in entries),
        'untitled': untitled, 'highlighted': sum(bool(e.get('isNew')) for e in entries),
    }
    return entries, record


def parse_sources(config, formal, files):
    loaded, parts, records = {}, {}, {}
    emoji = config.get('foodEmojiParagraph') or {}
    for key, path in files.items():
        expected = config['parts'].get(key, {}).get('sha256')
        if expected and base.digest(path) != expected:
            raise ValueError(f'{key}: 源文件哈希与配置不符')
        structure = match.inspect_docx_package(path)
        if not match.word_structure_safe(structure):
            raise ValueError(f'{key}: Word 含修订、表格或文本框，需人工处理')
        loaded[key] = load_part(path, key, config.get('proseParagraphs', {}).get(key),
                                emoji.get('index') if emoji.get('part') == key else None,
                                config.get('specialChapters', ()), config.get('titleParagraphs', {}).get(key),
                                config.get('translationParagraphs', {}).get(key), config.get('joinedTitles', {}).get(key))
        book = part_book(key)
        parts[key], records[key] = normalize_part(
            key, loaded[key], aliased(formal[book]['entries'], aliases_of(config, book)),
            config.get('styleCollections', ()), config.get('sourceSplits', ()))
        records[key]['structure'] = structure
        records[key]['specialChapters'] = loaded[key]['chapters']
        print(json.dumps({'parsed': key, 'count': len(parts[key])}, ensure_ascii=True), flush=True)
    return loaded, parts, records


# ---------------------------------------------------------------- assembly

def oc_signature(entries):
    return tuple(merge.fingerprint(e) for e in entries if special(e) == 'oc')


def block_items(items, rule_side):
    depth, digest = rule_side['depth'], rule_side['pathSha256']
    return [e for e in items if len(e['path']) >= depth and path_sha(e['path'][:depth]) == digest]


def drop_duplicate_block(raw_drop, raw_keep, entries, rule):
    """Drop a copied section after comparing both copies at raw-parse level.

    Normalized forms can differ legitimately: standalone role cards only merge
    where old parent evidence exists, which is the kept location.
    """
    dropped = [(e['title'], e['tags']) for e in block_items(raw_drop, rule['drop'])]
    kept = [(e['title'], e['tags']) for e in block_items(raw_keep, rule['keep'])]
    if dropped != kept or len(dropped) != rule['rawItems']:
        raise ValueError(f"重复段 {rule.get('note', '')} 两份不再逐字相同，停下复核")
    remove = {id(e) for e in block_items(entries, rule['drop'])}
    return [e for e in entries if id(e) not in remove]


def assemble(config, loaded, parts, formal):
    sources, stats = {}, {}
    for book in BOOKS:
        keys = ordered_parts(config['parts'], book)
        if not keys:
            raise ValueError(f'{book}: 配置里没有对应分卷')
        if book in REGULAR_BOOKS:
            if len(keys) != 1:
                raise ValueError(f'{book}: 常规本应只有一份 Word')
            sources[book] = parts[keys[0]]
            continue
        halves = [config['parts'][k]['half'] for k in keys]
        if halves != sorted(halves, key=lambda h: h != 'upper'):
            raise ValueError(f'{book}: 上半册的分卷必须排在下半册之前')
        first = keys[0]
        kept = {first: list(parts[first])}
        first_artists = Counter(merge.fingerprint(e) for e in parts[first] if special(e) == 'artist')
        oc_groups = [oc_signature(parts[first])]
        front = {}
        for key in keys[1:]:
            artists = Counter(merge.fingerprint(e) for e in parts[key] if special(e) == 'artist')
            if artists - first_artists:
                raise ValueError(f'{key}: 重复的画师组不是 {first} 的精确子集')
            oc = oc_signature(parts[key])
            drop_oc = oc in oc_groups
            if oc and not drop_oc:
                oc_groups.append(oc)
            kept[key] = [e for e in parts[key]
                         if not (special(e) == 'artist' or (special(e) == 'oc' and drop_oc))]
            front[key] = {'artistRemoved': sum(artists.values()), 'ocRemoved': len(oc) if drop_oc else 0,
                          'ocKeptAsOwnGroup': bool(oc) and not drop_oc}
        blocks = []
        for rule in config.get('duplicateBlocks', []):
            drop_key, keep_key = rule['drop']['part'], rule['keep']['part']
            if part_book(drop_key) != book:
                continue
            before = len(kept[drop_key])
            kept[drop_key] = drop_duplicate_block(loaded[drop_key]['raw'], loaded[keep_key]['raw'], kept[drop_key], rule)
            blocks.append({'note': rule.get('note', ''), 'removed': before - len(kept[drop_key])})
        for key in keys:
            leftover = [e['title'] for e in kept[key] if match.STANDALONE_ROLE_TITLE_RE.fullmatch(e['title'])]
            if leftover:
                raise ValueError(f'{key}: 还剩单独成段的角色卡 {leftover}，先查目录改名或父卡')
        upper = [e for k in keys if config['parts'][k]['half'] == 'upper' for e in kept[k]]
        lower = [e for k in keys if config['parts'][k]['half'] == 'lower' for e in kept[k]]
        corrections = []
        if book == 'suozhang_r18':
            formal_halves = merge.partition_formal_entries(formal[book]['entries'])
            for half in ('upper', 'lower'):
                group = upper if half == 'upper' else lower
                group, audit = merge.apply_known_source_title_corrections(group, formal_halves[half], half=half)
                corrections.extend(audit)
                if half == 'upper':
                    upper = group
                else:
                    lower = group
        merged = merge.merge_source_halves(upper, lower)
        sources[book] = merged['merged']
        stats[book] = {'frontMatter': front, 'duplicateBlocks': blocks, 'titleCorrections': corrections,
                       **merged['stats'], **merged['special']}
    own_chapter_titles(config, formal, sources)
    return sources, stats


def chapter_paths(config, book):
    return {(rule['chapter'],) for rule in config.get('specialChapters', []) if part_book(rule['part']) == book}


def own_chapter_titles(config, formal, sources):
    """Titles in a special chapter belong to the site, not the source.

    The source has no titles there; the site numbers them once and may later
    give them real names.  A card whose prompt is unchanged keeps the site's
    title; a new or edited one continues the numbering after the highest
    number in use.  On the first import (no old cards) the source numbering stays.
    """
    for rule in config.get('specialChapters', []):
        book = part_book(rule['part'])
        path = [rule['chapter']]
        olds = [e for e in formal[book]['entries'] if e.get('path') == path]
        if not olds:
            continue
        pool = defaultdict(list)
        for entry in olds:
            pool[body(entry)].append(entry['title'])
        totals_old = {text_sha: len(titles) for text_sha, titles in pool.items()}
        chapter_entries = [e for e in sources[book] if e.get('path') == path]
        # Adapted copies carry their own prompt; only the 原Tag cards are paired.
        counts_source = Counter(body(e) for e in chapter_entries
                                if not e['title'].endswith('：改编后'))
        consumed = Counter()
        pattern = re.compile(re.escape(rule['titlePrefix']) + r'（(\d+)）')
        numbers = [int(m.group(1)) for e in olds for m in [pattern.match(e['title'])] if m]
        next_number = max(numbers, default=0)
        main_title = None
        for entry in chapter_entries:
            adapted = entry['title'].endswith('：改编后')
            text_sha = body(entry)
            titles = pool.get(text_sha, [])
            pairable = (len(titles) > 1
                        and len(titles) + consumed[text_sha] == totals_old[text_sha] == counts_source[text_sha])
            if len(titles) == 1:
                entry['title'] = titles[0]
            elif pairable:
                # The author keeps repeated identical prompts (10.27: two copies
                # of one 藤蔓少女 group). Several old cards share that body too,
                # so give them back their titles in order instead of renumbering.
                entry['title'] = titles.pop(0)
                consumed[text_sha] += 1
            elif adapted and main_title:
                entry['title'] = f'{main_title}：改编后'
            else:
                next_number += 1
                entry['title'] = f"{rule['titlePrefix']}（{next_number}）"
            if not adapted:
                main_title = entry['title']


# ---------------------------------------------------------------- build

def classify_removals(book, removed, sources, pairs):
    sections = {e['path'][0] for e in sources[book] if e.get('path')}
    bodies = {other: {body(e) for e in sources[other]} for other in BOOKS if other != book}
    rows = []
    for entry in removed:
        first = alias_path(entry.get('path', []), pairs)[:1]
        targets = [other for other, found in bodies.items() if body(entry) in found]
        if first and first[0] not in sections:
            reason = 'sectionGone'
        elif targets:
            reason = 'movedTo'
        else:
            reason = 'unexplained'
        rows.append({'id': entry['id'], 'path': entry['path'], 'title': entry['title'],
                     'imaged': bool(entry.get('image')), 'reason': reason, 'targets': targets})
    return rows


def check_removals(config, book, rows):
    rules = config.get('removals', {}).get(book, {})
    retired = set(rules.get('retiredSections', []))
    moved = {target: set(ids) for target, ids in rules.get('movedTo', {}).items()}
    approved = set(rules.get('approvedIds', []))
    imaged_ok = set(rules.get('approvedImagedIds', []))
    for row in rows:
        if row['imaged'] and row['id'] not in imaged_ok:
            raise ValueError(f"带图词条将被删除：{row['id']}，必须由维护者确认后写进 approvedImagedIds")
        if row['reason'] == 'sectionGone' and row['path'][0] in retired:
            continue
        if row['reason'] == 'movedTo' and any(row['id'] in moved.get(t, ()) for t in row['targets']):
            continue
        if row['id'] in approved or row['id'] in imaged_ok:
            continue
        raise ValueError(f"未获批准的删除：{row['id']}（{row['reason']}）")
    removed_ids = {row['id'] for row in rows}
    for target, ids in moved.items():
        if ids - removed_ids:
            raise ValueError(f'{book}: movedTo 里有没被删除的 ID {sorted(ids - removed_ids)}')


def resolve_manual_matches(config, book, old_entries, new_entries):
    """Audited old ID → new card pairs, located by the new card's signature hash."""
    rules = config.get('manualMatches', {}).get(book, [])
    if not rules:
        return {}
    old_index = {e['id']: i for i, e in enumerate(old_entries)}
    by_signature = defaultdict(list)
    for i, entry in enumerate(new_entries):
        by_signature[match.override_signature_sha256(entry)].append(i)
    forced = {}
    for rule in rules:
        hits = by_signature.get(rule['newSignatureSha256'], [])
        if rule['oldId'] not in old_index or len(hits) != 1:
            raise ValueError(f"手工匹配 {rule['oldId']} 对不上了（源稿或旧数据变了），停下复核")
        forced[old_index[rule['oldId']]] = hits[0]
    return forced


def refresh_source_notes(book, built, sources, owned_paths):
    """Notes the author wrote in the source follow it on matched cards too.

    Built entries follow source order.  A translation or 设置 label in the
    source replaces the card's note; a special chapter owns its notes
    outright, elsewhere a site-added note stays when the source has none.
    Returns the IDs whose note came from the source.
    """
    if len(built) != len(sources) or any(entry.get('path') != source.get('path') for entry, source in zip(built, sources)):
        raise ValueError(f'{book}: 生成结果与源稿顺序对不上，备注无法跟随源稿')
    noted = set()
    for entry, source in zip(built, sources):
        if source.get('note'):
            entry['note'] = source['note']
            noted.add(entry['id'])
        elif tuple(entry.get('path', [])) in owned_paths:
            entry.pop('note', None)
    return noted


def build_books(config, formal, index, sources, *, reserved_cap=None, strict=True, enforce_removals=True):
    """Match and build all four books; ``strict=False`` reports instead of raising."""
    version = config['version']
    install_overrides(config)
    applied_books, results, problems = {}, {}, []
    for book in BOOKS:
        old = formal[book]
        entries = sources[book]
        pairs = aliases_of(config, book)
        try:
            source_overrides = match.apply_audited_source_new_overrides(entries, book, version)
        except ValueError as exc:
            if strict:
                raise
            problems.append(f'{book}: {exc}')
            source_overrides = []
        forced = resolve_manual_matches(config, book, old['entries'], entries)
        result = match.match_entries(aliased(old['entries'], pairs), entries, forced_pairs=forced)
        if result['review'] or result['validationIssues']:
            message = f"{book}: {len(result['review'])} 条待复核，校验问题 {result['validationIssues']}"
            if strict:
                raise ValueError(message)
            problems.append(message)
            results[book] = {'match': result}
            continue
        builder = merge.build_applied_codex if book == 'suozhang_r18' else match.build_applied_codex
        cap = id_caps(old['entries']) if reserved_cap else None
        applied, stats = builder(old, entries, result, version, reserved_ids=reserved_ids(book, cap),
                                 previous_update_filters=next(e for e in index if e['id'] == book).get('updateFilters'))
        for entry in applied['entries']:
            entry.pop('sourceHalf', None)
            entry.pop('sourceIndex', None)
        owned_paths = chapter_paths(config, book)
        noted = refresh_source_notes(book, applied['entries'], entries, owned_paths)
        if book == 'suozhang_r18':
            stats['auditedNewOverrideIds'] = match.apply_audited_new_overrides(applied['entries'], book, version)
        if sorted(source_overrides) != sorted(stats.get('auditedNewOverrideIds', [])):
            raise ValueError(f'{book}: NEW 修正没有落到配置里的稳定 ID 上')
        base.validate_book(applied)
        by_id = {e['id']: e for e in applied['entries']}
        for old_entry in old['entries']:
            new_entry = by_id.get(old_entry['id'])
            if new_entry is None:
                continue
            for key, value in old_entry.items():
                owned = {'note'} if old_entry['id'] in noted or tuple(old_entry.get('path', [])) in owned_paths else set()
                if key not in {*match.SOURCE_CONTENT_KEYS, 'updateBatches', 'sourceHalf', 'sourceIndex', *owned} \
                        and new_entry.get(key) != value:
                    raise ValueError(f"{book}: 继承字段漂移 {old_entry['id']} {key}")
        applied['title'] = base.TITLES[book]
        removed_ids = set(stats['removedIds'])
        removed = [e for e in old['entries'] if e['id'] in removed_ids]
        applied_books[book] = applied
        results[book] = {'match': result, 'stats': stats, 'removed': removed,
                         'summary': {'entryCount': applied['entryCount'], 'newIdCount': stats['newIdCount'],
                                     'removedCount': stats['removedCount'], 'imagedCount': applied['imagedCount'],
                                     'latestUpdateBatchCount': stats['latestUpdateBatchCount']}}
    for book, result in results.items():
        if 'removed' not in result:
            continue
        rows = classify_removals(book, result['removed'], sources, aliases_of(config, book))
        result['removals'] = rows
        if not enforce_removals:
            continue
        try:
            check_removals(config, book, rows)
        except ValueError as exc:
            if strict:
                raise
            problems.append(str(exc))
    new_index = index
    if len(applied_books) == len(BOOKS):
        old_index = {e['id']: e for e in index}
        for book in BOOKS:
            applied = applied_books[book]
            new_index = (merge.build_updated_codex_index(new_index, applied) if book == 'suozhang_r18'
                         else match.build_updated_codex_index(new_index, applied, book))
            meta = next(e for e in new_index if e['id'] == book)
            if meta['title'] != base.TITLES[book]:
                raise ValueError(f'{book}: 书目标题漂移')
            match.validate_update_batch_contract(applied, meta['updateFilters'])
        if any(item != next(e for e in new_index if e['id'] == item['id'])
               for item in old_index.values() if item['id'] not in BOOKS):
            raise ValueError('无关书目的索引被改动')
    return applied_books, results, new_index, problems


def replay_baseline(config, formal):
    base_dir = ROOT / config['baseline']['dir']
    summaries = {}
    for book, name in SNAPSHOTS.items():
        snapshot = base.read_json(base_dir / name)
        if snapshot.get('version') != config['baseVersion']:
            raise ValueError(f'{book}: 回放快照版本不是 {config["baseVersion"]}')
        owned = chapter_paths(config, book)

        def untitled(entries):  # site-owned titles in a special chapter are not drift
            return [dict(e, title='') if tuple(e.get('path', [])) in owned else e for e in entries]

        result = match.match_entries(untitled(formal[book]['entries']), untitled(snapshot['entries']))
        summaries[book] = {'summary': result['summary'],
                           'drift': [{'id': m['old']['id'], 'fields': sorted(f for f in m['changes'] if f != 'isNew')}
                                     for m in result['matches'] if any(f != 'isNew' for f in m['changes'])],
                           'unmatchedOld': [e['id'] for e in result['unmatchedOld']],
                           'unmatchedNew': len(result['unmatchedNew'])}
    return summaries


def replay_failures(config, baseline):
    """Books whose replay fails beyond the local edits listed in ``acceptedDrift``."""
    failed = []
    for book, item in baseline.items():
        if item['summary']['strictReplayPass']:
            continue
        allowed = {(row['id'], tuple(sorted(row['fields']))) for row in config.get('acceptedDrift', {}).get(book, [])}
        drift = {(row['id'], tuple(row['fields'])) for row in item['drift']}
        if item['unmatchedOld'] or item['unmatchedNew'] or drift != allowed:
            failed.append(book)
    return failed


def load_formal(formal_dir=None):
    folder = Path(formal_dir) if formal_dir else ROOT / 'site/data'
    formal = {book: base.read_json(folder / f'{book}.json') for book in BOOKS}
    index = base.read_json(folder / 'codexes.json')
    return formal, index


# ---------------------------------------------------------------- prepare / apply

def build_from_config(config, formal_dir=None):
    if config.get('historical'):
        raise ValueError('历史配置只供调研参照，不能 prepare')
    formal, index = load_formal(formal_dir)
    if any(book['version'] != config['baseVersion'] for book in formal.values()):
        raise ValueError(f"正式数据不是 {config['baseVersion']}，先确认上一版是否已应用")
    baseline = replay_baseline(config, formal)
    failed = replay_failures(config, baseline)
    if failed:
        raise ValueError(f'基线严格回放失败：{failed}；先查 survey 报告里的「本地修改漂移」，'
                         '维护者确认保留的写进 acceptedDrift')
    files = source_files(NEW_DATA / config['source']['dir'])
    if set(files) != set(config['parts']):
        raise ValueError(f'源文件分卷与配置不符：{sorted(files)} vs {sorted(config["parts"])}')
    loaded, parts, records = parse_sources(config, formal, files)
    sources, merges = assemble(config, loaded, parts, formal)
    applied_books, results, new_index, _ = build_books(config, formal, index, sources,
                                                      reserved_cap=formal_dir is not None, strict=True)
    return {'baseline': baseline, 'records': records, 'sources': sources, 'merges': merges,
            'applied': applied_books, 'results': results, 'index': new_index, 'loaded': loaded, 'formal': formal}


def prompt_text(entry):
    roles = [str(value) for prompt in entry.get('characterPrompts') or [] for value in prompt.values()]
    return '\n'.join([entry.get('tags') or '', *roles])


def prose_outcomes(config, built):
    """Where every title-like short N5 line ended up in this build.

    Registered lines should sit inside the card above; unregistered ones are
    either a card title or an empty title that forms no card. Lets a count
    difference be pinned to named cards instead of reasoning about it.
    """
    docs = {key: item['doc'] for key, item in built['loaded'].items()}
    rows = []
    for row in prose_candidates(docs, {}):
        key, text = row['part'], row['text']
        entries = [e for e in built['applied'][part_book(key)]['entries'] if e.get('path') == row['path']]
        registered = str(row['paragraph']) in config.get('proseParagraphs', {}).get(key, {})
        out = {'part': key, 'paragraph': row['paragraph'], 'text': text, 'registered': registered,
               'afterImageText': row['afterImageText']}
        if registered:
            host = next((e for e in entries if text in prompt_text(e)), None)
            out.update(outcome='merged' if host else 'missing', card=host and [host['id'], host['title']])
        else:
            spot = next((i for i, e in enumerate(entries) if e['title'] == text), None)
            if spot is None:
                out['outcome'] = 'noCard'
            else:
                card = entries[spot]
                out.update(outcome='ownCard', card=[card['id'], card['title']],
                           cardAbove=[entries[spot - 1]['id'], entries[spot - 1]['title']] if spot else None)
        rows.append(out)
    for key, registered in config.get('titleParagraphs', {}).items():
        lines = {index: (text, path) for index, text, path in body_lines(built['loaded'][key]['doc'])}
        for index in sorted(map(int, registered)):
            text, path = lines[index]
            card = next((e for e in built['applied'][part_book(key)]['entries']
                         if e.get('path') == path and e['title'] == text), None)
            rows.append({'part': key, 'paragraph': index, 'text': text, 'registered': True, 'afterImageText': False,
                         'outcome': 'titled' if card else 'titleMissing', 'card': card and [card['id'], card['title']]})
    return rows


def render_prose_outcomes(rows):
    lines = ['# 疑似续行的去向（本次预览）', '',
             '每一行短中文句在这次预览里落到了哪里，末尾几行是 `titleParagraphs` 登记成标题的行。'
             '条数和草稿差几条，就从这里数出是哪几张卡。', '',
             '| 分卷 | 段 | 原文 | 已登记 | 结果 |', '| --- | ---: | --- | --- | --- |']
    for row in rows:
        card = row.get('card')
        result = {
            'merged': lambda: f"并进 `{card[0]}`「{card[1]}」的正文",
            'missing': lambda: '**已登记，却没在这个目录的任何卡里找到这句——看原文**',
            'noCard': lambda: '空标题，不成卡（下面的正文归到后一个标题）',
            'ownCard': lambda: (f"**自成一卡** `{card[0]}`；上一张卡是 `{row['cardAbove'][0]}`「{row['cardAbove'][1]}」"
                                if row.get('cardAbove') else f"**自成一卡** `{card[0]}`"),
            'titled': lambda: f"登记为标题，成了卡 `{card[0]}`",
            'titleMissing': lambda: '**登记为标题，却没找到这个标题的卡——看原文**',
        }[row['outcome']]()
        hint = '（接在画面文字后面）' if row['afterImageText'] else ''
        lines.append(f"| {row['part']} | {row['paragraph']} | {row['text']}{hint} | {'是' if row['registered'] else '否'} | {result} |")
    return '\n'.join(lines) + '\n'


SHORT_CN_LINE_RE = re.compile(r'[^A-Za-z,:：]{2,30}')


def audit_flags(entry):
    """Shapes that misparsed cards had; a hint for the blind-spot check, not a verdict."""
    flags = []
    if UNTITLED_RE.search(entry['title']) and not entry.get('note'):  # a translated card is untitled by design
        flags.append('无标题')
    if sentence_title(entry['title']):
        flags.append('标题是长句')
    elif len(entry['title']) > 30 or re.search(r'[,，]', entry['title']):
        flags.append('标题像正文')
    numbers = [int(m.group(1)) for prompt in entry.get('characterPrompts') or []
               for m in [re.match(r'char(\d+)$', str(prompt.get('label', '')), re.I)] if m]
    if any(later == 1 and earlier > 1 for earlier, later in zip(numbers, numbers[1:])):
        flags.append('角色框编号重来')
    if any(SHORT_CN_LINE_RE.fullmatch(line.strip()) for line in prompt_text(entry).split('\n') if line.strip()):
        flags.append('正文里夹着中文短句')
    return flags


def changes_report(built):
    """Every new card and every old card whose content changed, per book."""
    report = {}
    for book in BOOKS:
        result = built['results'][book]['match']
        entries = built['applied'][book]['entries']
        old_ids = {m['old']['id'] for m in result['matches']}
        by_id = {e['id']: e for e in entries}
        old_notes = {e['id']: e.get('note') or '' for e in built.get('formal', {}).get(book, {}).get('entries', [])}

        def fields(m):  # the matcher does not compare notes
            note = old_notes.get(m['old']['id'], '') != (by_id[m['old']['id']].get('note') or '')
            return [f for f in m['changes'] if f != 'isNew'] + (['note'] if note else [])

        added = [{'id': e['id'], 'title': e['title'], 'path': e['path'], 'tags': e.get('tags', ''),
                  'characterPrompts': e.get('characterPrompts', []), 'flags': audit_flags(e)}
                 for e in entries if e['id'] not in old_ids]
        changed = [{'id': m['old']['id'], 'path': by_id[m['old']['id']]['path'], 'oldTitle': m['old']['title'],
                    'title': by_id[m['old']['id']]['title'],
                    'fields': fields(m), 'flags': audit_flags(by_id[m['old']['id']])}
                   for m in result['matches'] if fields(m)]
        report[book] = {'added': added, 'changed': changed}
    return report


def render_changes(report):
    lines = ['# 新增与改动（本次预览）', '',
             '给盲区检查用：每本列出全部新增卡和内容有改动的旧卡。⚑ 是工具按以往读错过的形状挑出来的，'
             '不一定有错，逐张对照 Word 原文。', '']
    for book, item in report.items():
        flagged = [row for row in item['added'] + item['changed'] if row['flags']]
        lines += [f"## {book}：新增 {len(item['added'])}，改动 {len(item['changed'])}，可疑 {len(flagged)}", '']
        if flagged:
            lines += ['### 可疑', '']
            for row in flagged:
                lines.append(f"- `{row['id']}` {' / '.join(row.get('path', []))} › {row['title']}　⚑ {'、'.join(row['flags'])}")
                if 'tags' in row:
                    roles = '、'.join(str(p.get('label', '')) for p in row['characterPrompts'])
                    lines.append(f"  正文：{row['tags'][:100]}{'…' if len(row['tags']) > 100 else ''}"
                                 + (f"｜角色框：{roles}" if roles else ''))
            lines.append('')
        lines += ['### 全部新增', '']
        lines += [f"- `{row['id']}` {' / '.join(row['path'])} › {row['title']}" for row in item['added']] or ['- 无']
        lines += ['', '### 内容有改动的旧卡', '']
        lines += [f"- `{row['id']}` {row['oldTitle'][:40]}{' → ' + row['title'][:40] if row['title'] != row['oldTitle'] else ''}"
                  f"（{'、'.join(row['fields'])}）" for row in item['changed']] or ['- 无']
        lines.append('')
    return '\n'.join(lines)


def preview(config, out, *, formal_dir=None):
    """Build exactly as prepare does, but only print the result; nothing is frozen."""
    built = build_from_config(config, formal_dir)
    summary = {}
    for book in BOOKS:
        result = built['results'][book]
        expected = config.get('expected', {}).get(book)
        summary[book] = {**result['summary'],
                         'expected': expected,
                         'matchesExpected': None if expected is None else
                         all(result['summary'][k] == v for k, v in expected.items()),
                         'removalReasons': dict(Counter(row['reason'] for row in result['removals'])),
                         'newOverrides': result['stats'].get('auditedNewOverrideIds', [])}
    outcomes = prose_outcomes(config, built)
    out.mkdir(parents=True, exist_ok=True)
    base.write_json(out / 'preview.json', summary)
    (out / '续行去向.md').write_text(render_prose_outcomes(outcomes), encoding='utf-8')
    changes = changes_report(built)
    base.write_json(out / 'changes.json', changes)
    (out / '新增与改动.md').write_text(render_changes(changes), encoding='utf-8')
    print(json.dumps(summary, ensure_ascii=False, indent=1))
    print(json.dumps({'proseLines': dict(Counter(row['outcome'] for row in outcomes)),
                      'detail': str(out / '续行去向.md')}, ensure_ascii=False))
    print(json.dumps({'flaggedCards': {book: sum(bool(row['flags']) for row in item['added'] + item['changed'])
                                       for book, item in changes.items()},
                      'detail': str(out / '新增与改动.md')}, ensure_ascii=False))
    off = [book for book, item in summary.items() if item['matchesExpected'] is False]
    if off:
        print(f'注意：{off} 的条数与配置里的 expected 不一致；差额要指到具体的卡，解释通了再改 expected')
    return summary


def prepare(config, out, *, formal_dir=None, write_snapshots=True):
    out.mkdir(parents=True, exist_ok=True)
    if (out / 'applied.json').exists():
        raise ValueError('这份计划已经应用过；要复验用 verify')
    before = base.data_hashes()
    built = build_from_config(config, formal_dir)
    baseline, records, sources, merges = built['baseline'], built['records'], built['sources'], built['merges']
    applied_books, results, new_index = built['applied'], built['results'], built['index']
    for book in BOOKS:
        actual = {key: results[book]['summary'][key] for key in ('entryCount', 'newIdCount', 'removedCount')}
        if actual != config['expected'][book]:
            raise ValueError(f'{book}: {actual} 与配置 expected {config["expected"][book]} 不符')
        base.write_json(out / 'candidates' / f'{book}.json', applied_books[book])
    base.write_json(out / 'candidates' / 'codexes.json', new_index)
    if write_snapshots:
        for book, name in SNAPSHOTS.items():
            payload = {'version': config['version'], 'entries': sources[book]}
            if book.endswith('_r18'):
                payload = {'codexId': book, 'version': config['version'], 'entryCount': len(sources[book]),
                           'entries': sources[book]}
            base.write_json(out / name, payload)
    payloads = {f'site/data/{book}.json': out / 'candidates' / f'{book}.json' for book in (*BOOKS, 'codexes')}
    if formal_dir is None and base.data_hashes() != before:
        raise ValueError('计划期间正式数据被改动；等其它会话结束后重跑')
    plan = {
        'version': config['version'], 'config': config['_path'], 'baseline': baseline, 'sources': records,
        'merges': merges, 'beforeDataHashes': before,
        'books': {book: {'summary': results[book]['summary'], 'match': base.compact_result(results[book]['match']),
                         'removals': results[book]['removals'],
                         'auditedNewOverrideIds': results[book]['stats'].get('auditedNewOverrideIds', [])}
                  for book in BOOKS},
        'targets': {rel: {'candidate': str(path), 'sha256': base.digest(path),
                          'beforeSha256': before.get(str(Path(rel)))} for rel, path in payloads.items()},
    }
    base.write_json(out / 'plan.json', plan)
    print(json.dumps({book: results[book]['summary'] for book in BOOKS}, ensure_ascii=True, indent=1))
    return plan, sources


# ---------------------------------------------------------------- survey

def relocate_lines(previous, docs, field, still=None, models=('N5',)):
    """Find last version's reviewed lines again by text hash; positions move.

    A line is only looked for in parts of the model it was registered for,
    and ``still(kind)`` must hold for how the parser would read it now (no
    registration is needed once the author's text reads right by itself).
    """
    known = {(part_model(part), text_sha) for part, rows in previous.get(field, {}).items() for text_sha in rows.values()}
    found, used = {}, set()
    for key, doc in docs.items():
        model = part_model(key)
        if model not in models:
            continue
        classify = convert.classify_nai5_line if model == 'N5' else convert.classify
        for index, paragraph in enumerate(doc.paragraphs):
            text_sha = sha256_text(paragraph.text)
            if (model, text_sha) in known and (still is None or still(classify(paragraph.text.strip()))):
                found.setdefault(key, {})[str(index)] = text_sha
                used.add((model, text_sha))
    return found, sorted(text_sha for _, text_sha in known - used)


def relocate_prose(previous, docs):
    return relocate_lines(previous, docs, 'proseParagraphs', lambda kind: kind != 'tag', models=('N4.5', 'N5'))


def relocate_titles(previous, docs):
    return relocate_lines(previous, docs, 'titleParagraphs', lambda kind: kind == 'tag')


def relocate_translations(previous, docs):
    return relocate_lines(previous, docs, 'translationParagraphs', models=('N4.5', 'N5'))


def relocate_joined_titles(previous, docs):
    """``joinedTitles`` rows keep their cut position; the paragraph is found by its text hash."""
    known = {row['sha256']: row['at'] for part in previous.get('joinedTitles', {}).values() for row in part.values()}
    found, used = {}, set()
    for key, doc in docs.items():
        for index, paragraph in enumerate(doc.paragraphs):
            text_sha = sha256_text(paragraph.text)
            if text_sha in known:
                found.setdefault(key, {})[str(index)] = {'sha256': text_sha, 'at': known[text_sha]}
                used.add(text_sha)
    return found, sorted(set(known) - used)


def find_food_emoji(docs):
    hits = []
    for key, doc in docs.items():
        if key != 'N4.5-regular':
            continue
        for index, paragraph in enumerate(doc.paragraphs):
            if len(convert.paragraph_blips(paragraph)) == 6:
                try:
                    base.restore_food_emoji(doc, index)
                except (ValueError, KeyError):
                    continue
                hits.append({'part': key, 'index': index})
    return hits


def after_image_text(paragraphs, index, run):
    """True when the prompt above this line opens on-image text (``text:``).

    Walks up through body lines and the lines of the same candidate run; a
    role line without ``text:``, a real title or a blank line ends the search.
    Only a hint for the report: the parser never uses it.
    """
    for above in range(index - 1, max(index - 30, -1), -1):
        line = paragraphs[above].text.strip()
        if not line or convert.outline_lvl(paragraphs[above]) is not None:
            return False
        if ON_IMAGE_TEXT_RE.search(line):
            return True
        if CHAR_LINE_RE.match(line):
            return False
        if above not in run and convert.classify_nai5_line(line) == 'title':
            return False
    return False


def body_lines(doc):
    """(index, text, path) of card lines: after the Word TOC, outside headings, front matter and skipped chapters."""
    cats, seen_toc = [None] * 4, False
    for index, paragraph in enumerate(doc.paragraphs):
        if paragraph.style.name.startswith('toc'):
            seen_toc = True
            continue
        text = paragraph.text.strip()
        level = convert.outline_lvl(paragraph)
        if level is not None and level <= 3 and text:
            cats[level] = text
            cats[level + 1:] = [None] * (3 - level)
            continue
        if not seen_toc or not text:
            continue
        path = [c for c in cats if c]
        if ('法典相关网站' in path or path[-1:] == ['编纂者常用画师组'] or convert.is_compiler_oc_path(path)
                or convert.should_skip_path(path)):
            continue
        yield index, text, path


def title_candidates(docs, known):
    """Titles N5 reads as prompt text: a short pure-Chinese line with a full-width comma.

    Any comma makes an N5 line prompt text, so such a title silently becomes
    the body of an untitled card or of the card above.  It must follow a blank
    line, or leave a bracket open (10.13 broke one title inside its brackets).
    """
    rows = []
    for key, doc in docs.items():
        if part_model(key) != 'N5':
            continue
        paragraphs = doc.paragraphs
        for index, text, path in body_lines(doc):
            unclosed = text.count('（') + text.count('(') > text.count('）') + text.count(')')
            if str(index) in known.get(key, {}) or not index or (paragraphs[index - 1].text.strip() and not unclosed):
                continue
            if '，' not in text or len(text) > 30 or re.search(r'[A-Za-z,:：]', text):
                continue
            if convert.classify_nai5_line(text) != 'tag':
                continue
            following = paragraphs[index + 1].text.strip() if index + 1 < len(paragraphs) else ''
            if following:
                rows.append({'part': key, 'paragraph': index, 'path': path, 'text': text, 'next': following[:60]})
    return rows


# Section headings inside one prompt, not card titles: 「## Pose」, 「first panel: …」.
SECTION_LABEL_RE = re.compile(r'#|(?:first|second|third|fourth|fifth|top|bottom|left|right|middle) panel\b', re.I)


def sentence_title(text):
    """A real title is never a long natural-language sentence: a translation,
    a description, or a title run together with its prompt.  Bracketed notes
    such as 【推荐画风：…】 do not count."""
    bare = re.sub(r'【[^】]*】|[（(][^）)]*[）)]', '', text)
    return len(bare) > 50 or '。' in bare


def long_title_candidates(docs, skip):
    """Lines the parser would take as a card title although they read as a sentence.

    Checked line by line as the parser reads them; ``skip`` holds
    (part, paragraph) pairs already registered or reported elsewhere.
    """
    rows = []
    for key, doc in docs.items():
        classify = convert.classify_nai5_line if part_model(key) == 'N5' else convert.classify
        paragraphs = doc.paragraphs
        for index, _, path in body_lines(doc):
            if (key, index) in skip or path[:1] == ['页尾']:
                continue
            lines = convert.visible_lines(paragraphs[index].text)
            for at, line in enumerate(lines):
                if classify(line) != 'title' or not (sentence_title(line) or SECTION_LABEL_RE.match(line)):
                    continue
                previous = lines[at - 1] if at else (paragraphs[index - 1].text.strip() if index else '')
                rows.append({'part': key, 'paragraph': index, 'path': path, 'text': line[:60], 'previous': previous[:60],
                             'afterPrompt': bool(previous) and classify(previous.splitlines()[-1]) == 'tag',
                             'section': bool(SECTION_LABEL_RE.match(line))})
    return rows


def dropped_lines(docs, skip):
    """English lines the parser reads as neither title nor prompt and silently drops.

    N4.5 keeps an English line only when it carries tag syntax (commas,
    brackets, weights); prose written as plain sentences was lost this way
    (10.13 vol. 3 card codex_8489ac52-1382 lost three paragraphs).
    """
    rows = []
    for key, doc in docs.items():
        classify = convert.classify_nai5_line if part_model(key) == 'N5' else convert.classify
        paragraphs = doc.paragraphs
        for index, _, path in body_lines(doc):
            if (key, index) in skip or path[:1] == ['页尾']:
                continue
            for line in convert.visible_lines(paragraphs[index].text):
                if classify(line) == 'note' and len(re.findall(r'[A-Za-z]', line)) >= 12:
                    rows.append({'part': key, 'paragraph': index, 'path': path, 'text': line[:60],
                                 'previous': paragraphs[index - 1].text.strip()[:60] if index else ''})
    return rows


def prose_candidates(docs, known):
    """Short title-like lines wedged inside a prompt: may be prose, needs a human."""
    rows = []
    for key, doc in docs.items():
        if part_model(key) != 'N5':
            continue
        paragraphs = doc.paragraphs
        run = {int(index) for index in known.get(key, {})}
        for index, text, path in body_lines(doc):
            if str(index) in known.get(key, {}):
                continue
            if convert.classify_nai5_line(text) != 'title' or text in LABEL_TITLES or VARIANT_TITLE_RE.match(text):
                continue
            previous = paragraphs[index - 1].text.strip() if index else ''
            following = paragraphs[index + 1].text.strip() if index + 1 < len(paragraphs) else ''
            if not previous or convert.classify_nai5_line(previous) != 'tag':
                continue
            if following and FRESH_PROMPT_RE.match(following) and not CHAR_LINE_RE.match(following):
                continue
            run.add(index)
            rows.append({'part': key, 'paragraph': index, 'path': path, 'text': text,
                         'previous': previous[:60], 'next': following[:60],
                         'afterImageText': after_image_text(paragraphs, index, run)})
            # Lines right after it that also read as titles belong to the same run.
            for later in range(index + 1, len(paragraphs)):
                line = paragraphs[later].text.strip()
                if not line or convert.classify_nai5_line(line) != 'title' or str(later) in known.get(key, {}):
                    break
                run.add(later)
                rows.append({'part': key, 'paragraph': later, 'path': path, 'text': line,
                             'previous': paragraphs[later - 1].text.strip()[:60],
                             'next': paragraphs[later + 1].text.strip()[:60] if later + 1 < len(paragraphs) else '',
                             'runStart': index, 'afterImageText': after_image_text(paragraphs, later, run)})
    return rows


def detect_aliases(old_entries, new_entries):
    old_entries = [e for e in old_entries if not special(e)]
    new_entries = [e for e in new_entries if not special(e)]
    pairs = []
    for depth in (1, 2, 3):
        old_view = aliased(old_entries, pairs)
        old_by = defaultdict(list)
        for entry in old_view:
            if len(entry['path']) >= depth:
                old_by[tuple(entry['path'][:depth])].append(body(entry))
        new_by = defaultdict(set)
        for entry in new_entries:
            if len(entry['path']) >= depth:
                new_by[tuple(entry['path'][:depth])].add(body(entry))
        claimed = set()
        for gone in sorted(set(old_by) - set(new_by)):
            scores = sorted(((sum(b in new_by[fresh] for b in old_by[gone]), fresh)
                             for fresh in set(new_by) - set(old_by) - claimed), reverse=True)
            if not scores or scores[0][0] == 0:
                continue
            best, fresh = scores[0]
            unique = len(scores) == 1 or scores[1][0] < best
            if unique and best * 2 >= len(old_by[gone]):
                pairs.append((gone, fresh))
                claimed.add(fresh)
    return pairs


def assign_halves(book, keys, parts, formal):
    if book in REGULAR_BOOKS:
        return {keys[0]: None}, {}
    if book != 'suozhang_r18':
        return {key: 'upper' if i == 0 else 'lower' for i, key in enumerate(keys)}, {}
    where = defaultdict(set)
    for entry in formal[book]['entries']:
        half = 'upper' if entry['id'].startswith(merge.UPPER_PREFIX) else 'lower'
        where[body(entry)].add(half)
    halves, evidence = {}, {}
    for key in keys:
        counts = Counter(h for e in parts[key] if not special(e) for h in where.get(body(e), ()))
        evidence[key] = dict(counts)
        halves[key] = counts.most_common(1)[0][0] if counts else None
    return halves, evidence


def detect_duplicate_blocks(loaded, keys, formal_entries, pairs):
    """Identical sections in two places of one book, compared on raw items."""
    groups = defaultdict(list)
    for key in keys:
        by_path = defaultdict(list)
        for entry in loaded[key]['raw']:
            if entry['path'][:1] == ['编纂者杂项']:
                continue
            by_path[tuple(entry['path'])].append((entry['title'], entry['tags']))
        for path, items in by_path.items():
            if len(items) >= 2:
                groups[tuple(items)].append((key, path))
    formal_paths = Counter(tuple(alias_path(e['path'], pairs)) for e in formal_entries)
    rules = []
    for items, places in groups.items():
        if len(places) != 2:
            continue
        ranked = sorted(places, key=lambda place: (-formal_paths[place[1]], keys.index(place[0])))
        keep, drop = ranked
        rules.append({
            'keep': {'part': keep[0], 'pathSha256': path_sha(keep[1]), 'depth': len(keep[1])},
            'drop': {'part': drop[0], 'pathSha256': path_sha(drop[1]), 'depth': len(drop[1])},
            'rawItems': len(items),
            'note': f'{drop[0]} 复制了 {keep[0]} 的一个子目录',
            '_keepPath': list(keep[1]), '_dropPath': list(drop[1]),
            '_keepFormalEntries': formal_paths[keep[1]], '_dropFormalEntries': formal_paths[drop[1]],
        })
    return rules


def detect_style_collections(parts):
    """Every section holding 负面/设置-style label cards becomes a fold rule."""
    rules, orphans = [], []
    for key, entries in parts.items():
        hosts, last = defaultdict(dict), {}
        for entry in entries:
            path = tuple(entry['path'])
            if entry['title'] in LABEL_TITLES:
                if path in last:
                    hosts[path].setdefault(last[path], []).append(entry['title'])
                else:
                    orphans.append({'part': key, 'path': list(path), 'title': entry['title']})
            else:
                last[path] = entry['title']
        rules.extend({'part': key, 'path': list(path), 'hosts': table} for path, table in hosts.items())
    return rules, orphans


def new_anomalies(book, entries, result):
    added = {item['index'] for item in result['additions']}
    rows = {'unhighlightedAdditions': [], 'highlightedUnchanged': []}
    for index in sorted(added):
        entry = entries[index]
        if entry.get('isNew'):
            continue
        previous = entries[index - 1] if index else None
        neighbour = bool(previous and previous['path'] == entry['path'] and previous.get('isNew')
                         and (index - 1) in added)
        rows['unhighlightedAdditions'].append({'index': index, 'title': entry['title'], 'path': entry['path'],
                                               'pinkNeighbour': neighbour})
    for item in result['matches']:
        entry = entries[item['new']['index']]
        if entry.get('isNew') and not any(field not in ('isNew', 'path') for field in item['changes']):
            rows['highlightedUnchanged'].append({'id': item['old']['id'], 'title': entry['title']})
    return rows


def copied_additions(entries, result):
    """New cards repeating, word for word, a card that an old ID already went to.

    The old ID follows the first copy in source order, which may be the
    author's fresh duplicate rather than the original position.
    """
    taken = {}
    for item in result['matches']:
        taken.setdefault(body(entries[item['new']['index']]), item)
    rows = []
    for item in result['additions']:
        twin = taken.get(body(entries[item['index']]))
        if twin:
            rows.append({'index': item['index'], 'title': item['title'], 'path': item['path'],
                         'oldId': twin['old']['id'], 'oldWentTo': twin['new']['index'],
                         'oldPath': twin['new']['path'], 'imaged': bool(twin['old'].get('image'))})
    return rows


def describe_copy(book, copy):
    where = f"{book}: 新增「{copy['title']}」（{' / '.join(copy['path'])}）与旧 {copy['oldId']}{'（带图）' if copy['imaged'] else ''}逐字相同"
    if copy['oldPath'] != copy['path']:
        return (f"{where}，旧编号留在「{' / '.join(copy['oldPath'])}」：作者在别的目录重复收了这一条，"
                '这张按新增收即可，过目确认')
    side = '前面' if copy['oldWentTo'] < copy['index'] else '后面'
    return (f'{where}：同一目录里出现了两份，旧编号和配图给了{side}那一份、这张另取新号——'
            f'看原文确认{side}那份是不是作者新复制的，是就报维护者')


def followed_moves(result):
    """Matched cards that changed section without a pathAliases rule."""
    moves = Counter((tuple(item['old']['path']), tuple(item['new']['path']))
                    for item in result['matches'] if item['old']['path'] != item['new']['path'])
    return [{'from': list(old), 'to': list(new), 'count': count} for (old, new), count in sorted(moves.items())]


def survey(source_dir, *, formal_dir=None, previous=None, out=None):
    files = source_files(source_dir)
    version = source_version([p.name for p in files.values()])
    previous = previous or previous_config(version)
    out = out or default_out(version)
    if (out / 'applied.json').exists():
        raise ValueError(f'{out} 是已经应用过的导入存档，调研不能写进去；演练或重跑请用 --out-dir 指到别处')
    out.mkdir(parents=True, exist_ok=True)
    formal, index = load_formal(formal_dir)
    capped = formal_dir is not None
    decisions, notes = [], []

    formal_versions = {book: data['version'] for book, data in formal.items()}
    if set(formal_versions.values()) != {previous['version']}:
        decisions.append(f'正式数据版本 {formal_versions} 不是上一版 {previous["version"]}：上一版可能没应用，先查清再继续')
    replay = replay_baseline({'baseline': {'dir': previous['outDir']}, 'baseVersion': previous['version'],
                              'specialChapters': previous.get('specialChapters', [])}, formal)
    for book, item in replay.items():
        if not item['summary']['strictReplayPass']:
            drift = '、'.join(f"{row['id']}（{','.join(row['fields'])}）" for row in item['drift'][:10])
            if len(item['drift']) > 10:
                drift += '…'
            decisions.append(f'{book} 基线回放不通过：{len(item["drift"])} 条本地修改漂移'
                             + (f'（{drift}）' if drift else '')
                             + f'、{len(item["unmatchedOld"])} 条旧 ID 对不上——上一版之后有人手改过数据，逐条确认要不要保留')

    structures = {key: match.inspect_docx_package(path) for key, path in files.items()}
    for key, structure in structures.items():
        if not match.word_structure_safe(structure):
            decisions.append(f'{key} 含修订痕迹、表格或文本框：{structure}，解析器不支持，需人工处理')
    docs = {key: Document(path) for key, path in files.items()}
    prose, missing_prose = relocate_prose(previous, docs)
    if missing_prose:
        notes.append(f'上一版的 {len(missing_prose)} 行已审续行这次没找到（作者可能删改了），无需处理')
    titles, missing_titles = relocate_titles(previous, docs)
    if missing_titles:
        notes.append(f'上一版的 {len(missing_titles)} 行已审标题这次没找到（作者可能删改了），无需处理')
    translations, missing_translations = relocate_translations(previous, docs)
    if missing_translations:
        notes.append(f'上一版的 {len(missing_translations)} 行已审翻译这次没找到（作者可能删改了），无需处理')
    joined, missing_joined = relocate_joined_titles(previous, docs)
    if missing_joined:
        decisions.append(f'上一版的 {len(missing_joined)} 处「标题连着提示词」这次没找到：作者可能已分行，也可能改了那一行，看原文确认')
    for key, indexes in prune_translations(docs, prose, titles, translations).items():
        decisions.append(f'{key} 第 {indexes} 段登记过翻译，但这次上面已不是提示词：草稿去掉了这几段的登记，'
                         '看原文后重新判断（是长句的会出现在「被当成标题的长句」里）')
    emoji_hits = find_food_emoji(docs)
    if len(emoji_hits) != 1:
        decisions.append(f'食物图标段落找到 {len(emoji_hits)} 处（应为 1 处）：{emoji_hits}')
    emoji = emoji_hits[0] if len(emoji_hits) == 1 else None
    chapters = []
    for rule in previous.get('specialChapters', []):
        try:
            if rule['part'] not in docs:
                raise ValueError('分卷不在')
            chapter_bounds(docs[rule['part']], rule['chapter'])
            chapters.append(dict(rule))
        except ValueError as exc:
            decisions.append(f"登记的特别章节「{rule['chapter']}」这次找不到（{exc}），草稿已去掉，核对原文后确认")
    candidates = prose_candidates(docs, prose)
    if candidates:
        hinted = sum(row['afterImageText'] for row in candidates)
        decisions.append(f'N5 有 {len(candidates)} 行疑似续行的短句被当成了标题'
                         f'{f"（其中 {hinted} 行接在 text: 画面文字后面）" if hinted else ""}，'
                         '逐条看原文后决定是否加进 proseParagraphs')
    reviewed = {(key, int(index)) for field in (prose, titles, translations, joined)
                for key, rows in field.items() for index in rows}
    reviewed |= {(row['part'], row['paragraph']) for row in candidates}
    sentence_rows = long_title_candidates(docs, reviewed)
    if sentence_rows:
        paired = sum(row['afterPrompt'] for row in sentence_rows)
        sections = sum(row['section'] for row in sentence_rows)
        decisions.append(f'{len(sentence_rows)} 行长句或正文分段被当成了标题（正常标题不会是一段自然语言；其中 {paired} 行紧跟在提示词下面，'
                         f'多半是那条提示词的中文翻译；{sections} 行像「## …」「first panel: …」这样的正文分段），'
                         '逐条看原文：翻译加进 translationParagraphs，正文加进 proseParagraphs')
    dropped_rows = dropped_lines(docs, reviewed)
    if dropped_rows:
        decisions.append(f'{len(dropped_rows)} 行英文被解析器当成说明直接丢掉了（不带逗号、括号的整句英文），'
                         '逐条看原文：是正文就加进 proseParagraphs')
    title_rows = title_candidates(docs, titles)
    if title_rows:
        decisions.append(f'N5 有 {len(title_rows)} 行疑似标题被当成了正文（带全角逗号的纯中文短行，在空行后或括号没闭合），'
                         '逐条看原文后决定是否加进 titleParagraphs')
    del docs

    draft = {
        'version': version, 'baseVersion': previous['version'],
        'source': {'dir': Path(os.path.abspath(source_dir)).relative_to(os.path.abspath(NEW_DATA)).as_posix()},
        'baseline': {'dir': previous['outDir']},
        'outDir': Path(os.path.abspath(out)).relative_to(os.path.abspath(ROOT)).as_posix(),
        'parts': {key: {'sha256': base.digest(path), 'half': None} for key, path in files.items()},
        'pathAliases': {}, 'proseParagraphs': prose, 'titleParagraphs': titles, 'translationParagraphs': translations,
        'joinedTitles': joined,
        'foodEmojiParagraph': emoji,
        'duplicateBlocks': [], 'styleCollections': [], 'removals': {}, 'newOverrides': {},
        'manualMatches': {}, 'acceptedDrift': {}, 'sourceSplits': [], 'specialChapters': chapters,
        'expected': {}, 'decisions': [],
    }
    loaded = {key: load_part(path, key, prose.get(key), emoji['index'] if emoji and emoji['part'] == key else None,
                             chapters, titles.get(key), translations.get(key), joined.get(key))
              for key, path in files.items()}
    for key, item in loaded.items():
        for name, stats in item['chapters'].items():
            notes.append(f"特别章节「{name}」：{stats['filled']} 组收录、{stats['adapted']} 组带改编版、"
                         f"{stats['emptyTemplates']} 个空模板不收")
    known_styles = {(rule['part'], tuple(rule['path'])) for rule in previous.get('styleCollections', [])}
    quick = {key: base.checked_normalize(item['raw'], [])[0] for key, item in loaded.items()}
    for rule in previous.get('sourceSplits', []):
        if rule['part'] in quick and source_split_hit(quick[rule['part']], rule) is not None:
            draft['sourceSplits'].append(dict(rule))
            notes.append(f"已登记的源稿问题仍在，继续拆分：{rule.get('note', '')}")
        else:
            decisions.append(f"已登记的源稿问题不再适用（{rule.get('note', '')}）：作者可能已修复或改了这段，"
                             '草稿已去掉这条记录，请核对原文后确认')
    rules, orphans = detect_style_collections(quick)
    draft['styleCollections'] = rules
    for rule in rules:
        if (rule['part'], tuple(rule['path'])) not in known_styles:
            decisions.append(f"{rule['part']} 「{' / '.join(rule['path'])}」出现负面/设置等小标签卡，"
                             f"草稿按「并回上一张卡」处理：{rule['hosts']}——需维护者确认")
    for orphan in orphans:
        decisions.append(f"{orphan['part']} 「{' / '.join(orphan['path'])}」的小标签「{orphan['title']}」找不到所属卡，需人工看原文")

    # Pass 1: normalize against unaliased formal data to find halves, renamed sections and copied blocks.
    splits = draft['sourceSplits']
    first_pass = {key: normalize_part(key, item, formal[part_book(key)]['entries'], rules, splits)[0]
                  for key, item in loaded.items()}
    halves_evidence = {}
    for book in BOOKS:
        keys = ordered_parts(files, book)
        if not keys:
            decisions.append(f'{book}: 这次没有对应的 Word')
            continue
        halves, evidence = assign_halves(book, keys, first_pass, formal)
        halves_evidence.update(evidence)
        for key, half in halves.items():
            draft['parts'][key]['half'] = half
            if half is None and book not in REGULAR_BOOKS:
                decisions.append(f'{key} 无法按旧内容判断属于上册还是下册号段，需人工指定 half')
        pairs = detect_aliases(formal[book]['entries'], [e for k in keys for e in first_pass[k]])
        if pairs:
            draft['pathAliases'][book] = [[list(old), list(new)] for old, new in pairs]
        for rule in detect_duplicate_blocks(loaded, keys, formal[book]['entries'], pairs):
            draft['duplicateBlocks'].append({k: v for k, v in rule.items() if not k.startswith('_')})
            decisions.append(f"{book}: {rule['drop']['part']}「{' / '.join(rule['_dropPath'])}」与 "
                             f"{rule['keep']['part']}「{' / '.join(rule['_keepPath'])}」逐字相同共 {rule['rawItems']} 项；"
                             f"草稿保留旧数据所在位置（旧条目 {rule['_keepFormalEntries']} 对 {rule['_dropFormalEntries']}）——需维护者确认")
    del first_pass

    # Pass 2: a real build of the draft. Removals and NEW gaps are read off a
    # first build without approvals, then the draft is rebuilt with them.
    preview_problems, anomalies, removal_rows, moves = [], {}, {}, {}
    try:
        parts = {key: normalize_part(key, item, aliased(formal[part_book(key)]['entries'],
                                                        aliases_of(draft, part_book(key))), rules, splits)[0]
                 for key, item in loaded.items()}
        sources, merges = assemble(draft, loaded, parts, formal)
        for stats in merges.values():
            for key, front in stats['frontMatter'].items():
                if front['ocKeptAsOwnGroup']:
                    notes.append(f'{key} 的编纂者 OC 与前面各卷不同，按独立一组保留（历史上下册 OC 就是两组），请核对')
            for correction in stats['titleCorrections']:
                notes.append(f"已知勘误 {correction['formalId']}：{correction['status']}")
        for book in BOOKS:
            review = match.match_entries(aliased(formal[book]['entries'], aliases_of(draft, book)), sources[book])['review']
            for item in review:
                ranked = sorted(item['candidates'], key=lambda c: -c['similarity']['score'])
                top = ranked[0]
                clear = len(ranked) == 1 or ranked[1]['similarity']['score'] < top['similarity']['score']
                entry = sources[book][item['new']['index']]
                if clear and top['old']['path'] == entry['path']:
                    draft.setdefault('manualMatches', {}).setdefault(book, []).append(
                        {'oldId': top['old']['id'], 'newSignatureSha256': match.override_signature_sha256(entry),
                         'note': '同目录改写，按旧编号继承'})
                    untitled = UNTITLED_RE.search(top['old']['title'])
                    decisions.append(f"{book}: 新卡「{entry['title']}」与旧 {top['old']['id']}「{top['old']['title']}」"
                                     f"（{'带图' if top['old'].get('image') else '无图'}）相似度 {top['similarity']['score']:.2f}，"
                                     '草稿按改写处理、继承旧编号——需维护者确认'
                                     + ('。**旧卡是无标题卡**：先按手册第 12 节查是不是被上方新卡吞并，是就改登记 sourceSplits'
                                        if untitled else ''))
                else:
                    decisions.append(f"{book}: 新卡「{entry['title']}」有多个相近旧词条，需人工指定 manualMatches")
        for book in BOOKS:
            match.AUDITED_NEW_OVERRIDES.pop((book, version), None)
        applied, results, _, problems = build_books(dict(draft, newOverrides={}), formal, index, sources,
                                                    reserved_cap=capped, strict=False, enforce_removals=False)
        preview_problems.extend(problems)
        formal_by_id = {e['id']: e for book in BOOKS for e in formal[book]['entries']}
        moved_bodies = set()
        for book, result in results.items():
            rows = result.get('removals', [])
            removal_rows[book] = rows
            if not rows:
                continue
            retired = sorted({r['path'][0] for r in rows if r['reason'] == 'sectionGone'})
            moved = defaultdict(list)
            for row in rows:
                if row['reason'] == 'movedTo':
                    moved_bodies.add(body(formal_by_id[row['id']]))
                    for target in row['targets']:
                        moved[target].append(row['id'])
            draft['removals'][book] = {'retiredSections': retired, 'movedTo': dict(moved)}
            counts = dict(Counter(r['reason'] for r in rows))
            decisions.append(f'{book}: 将删除 {len(rows)} 条 {counts}——是否跟随源稿删除，需维护者确认')
            unexplained = [r['id'] for r in rows if r['reason'] == 'unexplained']
            imaged = [r['id'] for r in rows if r['imaged']]
            if unexplained:
                decisions.append(f'{book}: {len(unexplained)} 条删除原因不明，核实后写进 removals.approvedIds')
            if imaged:
                decisions.append(f'{book}: {len(imaged)} 条**带图**词条将被删除，必须维护者逐条确认后写进 approvedImagedIds')
        proposed = 0
        for book, result in results.items():
            if 'match' not in result:
                continue
            anomalies[book] = rows = new_anomalies(book, sources[book], result['match'])
            for item in result['match']['additions']:
                if UNTITLED_RE.search(item['title']):
                    decisions.append(f"{book}: 新增了无标题卡「{item['title']}」——标题可能被当成了正文，"
                                     '或这段被上方新卡隔开了；看原文')
            rows['copies'] = copied_additions(sources[book], result['match'])
            for copy in rows['copies']:
                (notes if copy['oldPath'] != copy['path'] else decisions).append(describe_copy(book, copy))
            moves[book] = followed_moves(result['match'])
            if moves[book]:
                notes.append(f"{book}: {sum(m['count'] for m in moves[book])} 条旧词条按内容跟到了别的目录"
                             f"（{len(moves[book])} 处，没有生成改名规则），见报告「目录改名」一节，过目即可")
            ids = {match.override_signature_sha256(e): e['id'] for e in applied.get(book, {}).get('entries', [])}
            for row in rows['unhighlightedAdditions']:
                entry = sources[book][row['index']]
                row['movedIn'] = body(entry) in moved_bodies
                digest = match.override_signature_sha256(entry)
                if row['pinkNeighbour'] and not row['movedIn'] and digest in ids:
                    draft['newOverrides'].setdefault(book, {})[ids[digest]] = {'signatureSha256': digest, 'isNew': True}
                    proposed += 1
        if proposed:
            decisions.append(f'没标粉、但紧挨着粉色新卡的真新增 {proposed} 条，草稿按 8.14 先例补进 NEW——需维护者确认')
        if not preview_problems:
            for book in BOOKS:
                match.AUDITED_NEW_OVERRIDES.pop((book, version), None)
            _, results, _, problems = build_books(draft, formal, index, sources, reserved_cap=capped, strict=False)
            preview_problems.extend(problems)
            for book, result in results.items():
                if 'summary' in result:
                    draft['expected'][book] = {k: result['summary'][k] for k in ('entryCount', 'newIdCount', 'removedCount')}
    except ValueError as exc:
        preview_problems.append(str(exc))
    for problem in preview_problems:
        decisions.append(f'预览受阻：{problem}')

    report = {'version': version, 'previous': previous['version'], 'files': {k: str(v) for k, v in files.items()},
              'structures': structures, 'baselineReplay': replay, 'halvesEvidence': halves_evidence,
              'removals': removal_rows, 'newAnomalies': anomalies, 'followedMoves': moves, 'proseCandidates': candidates,
              'titleCandidates': title_rows, 'sentenceTitles': sentence_rows, 'droppedLines': dropped_rows,
              'decisions': decisions, 'notes': notes}
    base.write_json(out / 'survey.json', report)
    base.write_json(out / 'draft-config.json', draft)
    (out / '调研报告.md').write_text(render_survey(report, draft), encoding='utf-8')
    print(json.dumps({'version': version, 'decisions': len(decisions), 'out': str(out)}, ensure_ascii=True))
    return report, draft


def render_survey(report, draft):
    lines = [f"# 所长 {report['version']} 调研报告", '',
             f"上一版：{report['previous']}。本报告由 `tools/import_suozhang.py survey` 生成，只读，没有写任何正式数据。", '',
             '## 需要维护者拍板 / 需要人工判断', '']
    lines += [f'- {item}' for item in report['decisions']] or ['- 无']
    lines += ['', '## 备注', ''] + ([f'- {item}' for item in report['notes']] or ['- 无'])
    lines += ['', '## 预计结果（草稿配置）', '', '| 法典 | 最终条数 | 新增 | 删除 |', '| --- | ---: | ---: | ---: |']
    for book, value in draft['expected'].items():
        lines.append(f"| `{book}` | {value['entryCount']} | {value['newIdCount']} | {value['removedCount']} |")
    lines += ['', '## 分卷与号段', '']
    for key, part in draft['parts'].items():
        lines.append(f"- {key}：half={part['half']}；旧内容命中 {report['halvesEvidence'].get(key, {})}")
    lines += ['', '## 目录改名（只用于匹配）', '']
    for book, pairs in draft['pathAliases'].items():
        lines += [f"- {book}：{' / '.join(old)} → {' / '.join(new)}" for old, new in pairs]
    followed = [(book, move) for book, rows in report.get('followedMoves', {}).items() for move in rows]
    if followed:
        lines += ['', '按内容自动跟到新目录的旧词条（不需要改名规则，编号和配图都在；过目即可）：', '']
        lines += [f"- {book}：{' / '.join(move['from'])} → {' / '.join(move['to'])}（{move['count']} 条）"
                  for book, move in followed]
    lines += ['', '## 删除明细', '']
    removed = [(book, row) for book, rows in report['removals'].items() for row in rows]
    for book, row in removed:
        lines.append(f"- {book} `{row['id']}` {' / '.join(row['path'])} › {row['title']}"
                     f"（{row['reason']}{'，带图' if row['imaged'] else ''}{('→' + ','.join(row['targets'])) if row['targets'] else ''}）")
    if not removed:
        lines.append('- 无：这次四本都没有删除')
    lines += ['', '## 更新标记异常', '']
    for book, rows in report['newAnomalies'].items():
        for row in rows['unhighlightedAdditions']:
            flag = '迁入' if row.get('movedIn') else ('紧挨粉色新卡' if row['pinkNeighbour'] else '需人工看')
            lines.append(f"- {book} 新增未标粉：{' / '.join(row['path'])} › {row['title']}（{flag}）")
        for row in rows['highlightedUnchanged']:
            lines.append(f"- {book} 旧词条内容没变却标粉：`{row['id']}` {row['title']}（按惯例跟随源稿）")
        for row in rows.get('copies', []):
            lines.append(f"- {describe_copy(book, row)}")
    lines += ['', '## 疑似续行短句', '']
    lines += ['标了「接在画面文字后面」的：前文有 `text:`，这句多半是画面上文字的下一行（续行）。只是提示，仍要看原文。', '']
    for row in report['proseCandidates']:
        hint = '｜**接在画面文字后面**' if row.get('afterImageText') else ''
        lines.append(f"- {row['part']} 第 {row['paragraph']} 段「{row['text']}」{hint}｜上一行：{row['previous']}｜下一行：{row['next']}")
    lines += ['', '## 疑似被当成正文的标题', '',
              'N5 只要一行里有逗号就当正文读，带全角逗号的中文标题会被并进上一张卡，或变成无标题卡的正文。'
              '是标题就把段落号和该段原文的 sha256 加进 `titleParagraphs`。', '']
    lines += [f"- {row['part']} 第 {row['paragraph']} 段「{row['text']}」｜下一行：{row['next']}"
              for row in report.get('titleCandidates', [])] or ['- 无']
    lines += ['', '## 被当成标题的长句', '',
              '正常标题不会是一段自然语言。紧跟在提示词下面的长句多半是那条提示词的中文翻译：'
              '把它的段落号和原文 sha256 加进 `translationParagraphs`，两段合成一张无标题卡，翻译进备注。'
              '也可能是标题和提示词写在了同一行，那种报维护者。', '']
    lines += [f"- {row['part']} 第 {row['paragraph']} 段「{row['text']}」{'｜**紧跟在提示词下面**' if row['afterPrompt'] else ''}"
              f"{'｜**像正文分段**' if row.get('section') else ''}｜上一行：{row['previous']}"
              for row in report.get('sentenceTitles', [])] or ['- 无']
    lines += ['', '## 被丢掉的英文行', '',
              'N4.5 只认带逗号、括号等 tag 符号的英文行；整句写的英文会被当成说明丢掉。是卡片正文就把段落号和原文 sha256 加进 `proseParagraphs`。', '']
    lines += [f"- {row['part']} 第 {row['paragraph']} 段「{row['text']}」｜上一行：{row['previous']}"
              for row in report.get('droppedLines', [])] or ['- 无']
    lines += ['', '下一步：把 `draft-config.json` 按拍板结果改好，另存为 `tools/suozhang_versions/<版本>.json`，再运行 prepare。', '']
    return '\n'.join(lines)


# ---------------------------------------------------------------- selfcheck

SELFCHECK_VERSION = '2026.9.25'


def selfcheck(skip_survey=False):
    config = load_config(config_path(SELFCHECK_VERSION))
    archive = ROOT / config['outDir']
    fixture = archive / 'before'
    scratch = ROOT / 'output' / 'suozhang-selfcheck'
    for needed in (fixture, archive / 'candidates', NEW_DATA / config['source']['dir']):
        if not needed.exists():
            print(f'SKIP：缺少存档 {needed}')
            return 0
    failures = []
    if not skip_survey:
        _, draft = survey(NEW_DATA / config['source']['dir'], formal_dir=fixture,
                          previous=load_config(config_path(config['baseVersion'])), out=scratch / 'survey')
        for key in ('parts', 'proseParagraphs', 'foodEmojiParagraph', 'styleCollections', 'removals',
                    'newOverrides', 'expected', 'pathAliases'):
            if draft[key] != config[key]:
                failures.append(f'survey 草稿的 {key} 与 {SELFCHECK_VERSION} 配置不同')
        if [{k: v for k, v in r.items() if k != 'note'} for r in draft['duplicateBlocks']] != \
                [{k: v for k, v in r.items() if k != 'note'} for r in config['duplicateBlocks']]:
            failures.append('survey 草稿的 duplicateBlocks 与配置不同')
    for book in BOOKS:
        match.AUDITED_NEW_OVERRIDES.pop((book, config['version']), None)
    run = scratch / 'prepare'
    for stale in (run / 'applied.json',):
        stale.unlink(missing_ok=True)
    _, sources = prepare(config, run, formal_dir=fixture, write_snapshots=False)
    for book, name in SNAPSHOTS.items():
        if json.loads(json.dumps(sources[book], ensure_ascii=False)) != base.read_json(archive / name)['entries']:
            failures.append(f'{book} 的合并源与存档快照不同')
    # Pictures added to new IDs after the archived run are environment, not logic.
    imaged = {}
    for book in BOOKS:
        old_ids = {e['id'] for e in base.read_json(fixture / f'{book}.json')['entries']}
        built = base.read_json(run / 'candidates' / f'{book}.json')
        for entry in built['entries']:
            if entry['id'] not in old_ids:
                for key in match.ASSET_KEYS:
                    entry.pop(key, None)
                entry['image'] = None
        built['imagedCount'] = imaged[book] = sum(bool(e.get('image')) for e in built['entries'])
        if built != base.read_json(archive / 'candidates' / f'{book}.json'):
            failures.append(f'{book} 候选与存档不同')
    index = base.read_json(run / 'candidates' / 'codexes.json')
    for item in index:
        if item['id'] in imaged:
            item['imagedCount'] = imaged[item['id']]
    if index != base.read_json(archive / 'candidates' / 'codexes.json'):
        failures.append('codexes 候选与存档不同')
    print(json.dumps({'selfcheck': 'FAIL' if failures else 'PASS', 'failures': failures}, ensure_ascii=False, indent=1))
    return 1 if failures else 0


# ---------------------------------------------------------------- compare

def entry_signature(entry):
    return json.dumps([entry.get('id'), entry.get('title'), entry.get('path'), body(entry), entry.get('note') or '',
                       bool(entry.get('isNew')), sorted(entry.get('updateBatches') or [])], ensure_ascii=False)


def content_signature(entry):
    """The card as a reader sees it, without its ID: survives renumbering."""
    return json.dumps([entry.get('title'), entry.get('path'), body(entry), entry.get('note') or ''], ensure_ascii=False)


def compare(config, reference, formal_dir, out):
    """Build two configs on the same base data and diff the four books entry by entry.

    ``onlyIn*`` compare whole entries, IDs included, so one extra card shifts
    every later new ID; ``contentOnlyIn*`` ignore IDs and name the cards that
    really differ.
    """
    built = {}
    for name, cfg in (('config', config), ('reference', reference)):
        for book in BOOKS:  # each config installs its own NEW table
            match.AUDITED_NEW_OVERRIDES.pop((book, cfg['version']), None)
        built[name] = build_from_config(cfg, formal_dir)
    report, passed = {}, True
    for book in BOOKS:
        mine = Counter(entry_signature(e) for e in built['config']['applied'][book]['entries'])
        theirs = Counter(entry_signature(e) for e in built['reference']['applied'][book]['entries'])
        only_mine, only_theirs = mine - theirs, theirs - mine

        def examples(counter):
            return [json.loads(sig)[:3] for sig in list(counter)[:12]]

        content_mine = Counter(content_signature(e) for e in built['config']['applied'][book]['entries'])
        content_theirs = Counter(content_signature(e) for e in built['reference']['applied'][book]['entries'])

        def cards(counter):
            return [[title, path, prompt[:80]] for title, path, prompt, _ in map(json.loads, counter)]

        report[book] = {'config': built['config']['results'][book]['summary'],
                        'reference': built['reference']['results'][book]['summary'],
                        'onlyInConfig': sum(only_mine.values()), 'onlyInReference': sum(only_theirs.values()),
                        'examplesOnlyInConfig': examples(only_mine), 'examplesOnlyInReference': examples(only_theirs),
                        'contentOnlyInConfig': cards(content_mine - content_theirs),
                        'contentOnlyInReference': cards(content_theirs - content_mine)}
        passed = passed and not only_mine and not only_theirs
    result = {'result': 'PASS' if passed else 'FAIL', 'books': report}
    out.mkdir(parents=True, exist_ok=True)
    base.write_json(out / 'compare.json', result)
    print(json.dumps({'result': result['result'],
                      'books': {b: {'onlyInConfig': v['onlyInConfig'], 'onlyInReference': v['onlyInReference'],
                                    'contentOnlyInConfig': len(v['contentOnlyInConfig']),
                                    'contentOnlyInReference': len(v['contentOnlyInReference'])}
                                for b, v in report.items()},
                      'report': str(out / 'compare.json')}, ensure_ascii=False, indent=1))
    return 0 if passed else 1


# ---------------------------------------------------------------- cli

def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = parser.add_subparsers(dest='command', required=True)
    p = sub.add_parser('survey', help='只读调研新版 Word')
    p.add_argument('--source-dir', type=Path, required=True, help='解压后放 docx 的目录')
    p.add_argument('--previous', type=Path, help='上一版配置；默认取 tools/suozhang_versions/ 里最近的一份')
    p.add_argument('--out-dir', type=Path)
    p.add_argument('--formal-dir', type=Path, help='演练用：拿这个目录里的四本当「导入前的正式数据」，只读')
    for name, help_text in (('preview', '按配置预览结果（只读，不冻结）'), ('prepare', '生成冻结候选与计划'),
                            ('apply', '应用冻结计划'), ('verify', '只读复验')):
        p = sub.add_parser(name, help=help_text)
        p.add_argument('--config', type=Path, required=True)
        p.add_argument('--out-dir', type=Path, help='默认取配置里的 outDir')
        if name == 'preview':
            p.add_argument('--formal-dir', type=Path, help='演练用：导入前的正式数据目录，只读')
    p = sub.add_parser('compare', help='两份配置在同一份底数据上重建并逐条比对（演练对答案；平时查改配置前后差在哪几张卡）')
    p.add_argument('--config', type=Path, required=True, help='要检验的配置')
    p.add_argument('--reference', type=Path, required=True, help='对照配置：标准答案，或改动前的配置')
    p.add_argument('--formal-dir', type=Path, help='导入前的正式数据目录，只读；默认 site/data')
    p.add_argument('--out-dir', type=Path, required=True)
    p = sub.add_parser('selfcheck', help='用 9.25 存档重放，确认工具结果不变')
    p.add_argument('--skip-survey', action='store_true')
    args = parser.parse_args()
    if args.command == 'survey':
        previous = load_config(args.previous) if args.previous else None
        survey(args.source_dir, previous=previous, out=args.out_dir, formal_dir=args.formal_dir)
        return 0
    if args.command == 'selfcheck':
        return selfcheck(args.skip_survey)
    if args.command == 'compare':
        return compare(load_config(args.config), load_config(args.reference), args.formal_dir, args.out_dir)
    config = load_config(args.config)
    out = args.out_dir or ROOT / config['outDir']
    if args.command == 'preview':
        if args.out_dir is None and (out / 'applied.json').exists():
            raise SystemExit(f'{out} 是已应用的导入存档；预览请用 --out-dir 指到别处')
        preview(config, out, formal_dir=args.formal_dir)
    elif args.command == 'prepare':
        prepare(config, out)
    elif args.command == 'apply':
        base.apply_plan(out)
    else:
        base.verify(out)
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
