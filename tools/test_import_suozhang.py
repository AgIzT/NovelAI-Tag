import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parent))
import codex_update_match as match
import import_suozhang as imp
import suozhang_r18_merge_match as merge


def card(title, path, tags, **extra):
    return {'title': title, 'path': list(path), 'tags': tags, 'isNew': False, **extra}


ARTIST = ['编纂者杂项', '编纂者常用画师组']
OC = ['编纂者杂项', '编纂者OC']


class NamingTests(unittest.TestCase):
    def test_part_names(self):
        self.assertEqual(imp.part_key('N4.5所长色色NovelAI个人法典（卷三）（2026.9.25版）.docx'), 'N4.5-v3')
        self.assertEqual(imp.part_key('N4.5所长色色NovelAI个人法典（卷十二）（2026.12.1版）.docx'), 'N4.5-v12')
        self.assertEqual(imp.part_key('N5所长色色NovelAI个人法典（下）（2026.9.25版）.docx'), 'N5-lower')
        self.assertEqual(imp.part_key('N5所长常规NovelAI个人法典（2026.9.25版）.docx'), 'N5-regular')
        with self.assertRaises(ValueError):
            imp.part_key('所长色色NovelAI个人法典（卷一）.docx')

    def test_word_lock_files_are_ignored(self):
        with tempfile.TemporaryDirectory() as folder:
            for name in ('N5所长常规NovelAI个人法典（2026.10.13版）.docx', '~$所长常规NovelAI个人法典（2026.10.13版）.docx'):
                (Path(folder) / name).write_bytes(b'')
            self.assertEqual(list(imp.source_files(folder)), ['N5-regular'])

    def test_version_comes_from_file_names(self):
        names = ['N5所长常规NovelAI个人法典（2026.10.13版，一般所长整理）.docx',
                 'N4.5所长常规NovelAI个人法典（2026.10.13版，一般所长整理）.docx']
        self.assertEqual(imp.source_version(names), '2026.10.13')
        self.assertEqual(imp.version_slug('2026.10.3'), '20261003')
        with self.assertRaises(ValueError):
            imp.source_version(names + ['N5所长色色（上）（2026.9.25版）.docx'])

    def test_volume_kinds_cannot_mix(self):
        with self.assertRaises(ValueError):
            imp.ordered_parts({'N4.5-v1': {}, 'N4.5-upper': {}}, 'suozhang_r18')
        self.assertEqual(imp.ordered_parts({'N4.5-v10': {}, 'N4.5-v2': {}}, 'suozhang_r18'), ['N4.5-v2', 'N4.5-v10'])


@mock.patch.object(merge, 'KNOWN_SOURCE_TITLE_CORRECTIONS', ())
class AssembleTests(unittest.TestCase):
    def config(self, parts):
        return {'parts': parts, 'duplicateBlocks': []}

    def test_repeated_front_matter_is_removed_only_when_proven(self):
        regular = {'N4.5-regular': [card('风景', ['场景'], 'sky,')], 'N5-regular': [card('风景', ['场景'], 'sky,')],
                   'N5-upper': [card('甲', ['场景'], 'a,')], 'N5-lower': [card('乙', ['场景'], 'b,')]}
        v1 = [card('NAI5时期', ARTIST, 'artist:a,', negative='logo,'), card('NAI4.5时期：2', ARTIST, 'artist:b,'),
              card('编纂者OC(1)', OC, 'girl,'), card('林间', ['场景'], 'walk,')]
        v2 = [card('NAI4.5时期：2', ARTIST, 'artist:b,'), card('编纂者OC(1)', OC, 'girl,'), card('夜景', ['场景'], 'night,')]
        v3 = [card('NAI4.5时期：2', ARTIST, 'artist:b,'), card('编纂者OC(1)', ['编纂者杂项', '编纂者oc二则'], 'boy,'),
              card('雨天', ['场景'], 'rain,')]
        parts = {**regular, 'N4.5-v1': v1, 'N4.5-v2': v2, 'N4.5-v3': v3}
        config = self.config({'N4.5-regular': {'half': None}, 'N5-regular': {'half': None},
                              'N5-upper': {'half': 'upper'}, 'N5-lower': {'half': 'lower'},
                              'N4.5-v1': {'half': 'upper'}, 'N4.5-v2': {'half': 'upper'}, 'N4.5-v3': {'half': 'lower'}})
        formal = {'suozhang_r18': {'entries': []}}
        sources, stats = imp.assemble(config, {}, parts, formal)
        titles = [e['title'] for e in sources['suozhang_r18']]
        self.assertEqual(titles, ['NAI5时期', 'NAI4.5时期：2', '编纂者OC(1)', '林间', '夜景', '编纂者OC(1)', '雨天'])
        self.assertTrue(stats['suozhang_r18']['frontMatter']['N4.5-v3']['ocKeptAsOwnGroup'])

        changed = {**parts, 'N4.5-v2': [card('NAI5时期', ARTIST, 'artist:a,', negative='text,')]}
        with self.assertRaises(ValueError):
            imp.assemble(config, {}, changed, formal)

    def test_upper_volumes_must_come_first(self):
        parts = {'N4.5-regular': [], 'N5-regular': [], 'N5-upper': [], 'N5-lower': [],
                 'N4.5-v1': [], 'N4.5-v2': []}
        config = self.config({'N4.5-regular': {'half': None}, 'N5-regular': {'half': None},
                              'N5-upper': {'half': 'upper'}, 'N5-lower': {'half': 'lower'},
                              'N4.5-v1': {'half': 'lower'}, 'N4.5-v2': {'half': 'upper'}})
        with self.assertRaises(ValueError):
            imp.assemble(config, {}, parts, {'suozhang_r18': {'entries': []}})


class DuplicateBlockTests(unittest.TestCase):
    def rule(self, items=2):
        return {'drop': {'part': 'v2', 'pathSha256': imp.path_sha(['甲', '多人']), 'depth': 2},
                'keep': {'part': 'v3', 'pathSha256': imp.path_sha(['乙', '多人']), 'depth': 2}, 'rawItems': items}

    def test_verbatim_copy_is_dropped_from_the_copy_location(self):
        keep = [card('一', ['乙', '多人'], 'a,'), card('char1', ['乙', '多人'], 'girl,')]
        drop = [card('别处', ['甲', '其它'], 'b,'), card('一', ['甲', '多人'], 'a,'), card('char1', ['甲', '多人'], 'girl,')]
        kept = imp.drop_duplicate_block(drop, keep, drop, self.rule())
        self.assertEqual([e['title'] for e in kept], ['别处'])

    def test_diverged_copy_blocks(self):
        keep = [card('一', ['乙', '多人'], 'a,'), card('二', ['乙', '多人'], 'c,')]
        drop = [card('一', ['甲', '多人'], 'a,'), card('二', ['甲', '多人'], 'changed,')]
        with self.assertRaises(ValueError):
            imp.drop_duplicate_block(drop, keep, drop, self.rule())

    def test_detection_keeps_the_historical_location(self):
        raw = {'v2': [card('一', ['甲', '多人'], 'a,'), card('二', ['甲', '多人'], 'b,')],
               'v3': [card('一', ['乙', '多人'], 'a,'), card('二', ['乙', '多人'], 'b,'), card('三', ['乙', '其它'], 'c,')]}
        loaded = {key: {'raw': items} for key, items in raw.items()}
        formal = [{'id': 'x-1', 'path': ['乙', '多人'], 'title': '一', 'tags': 'a,'}]
        rules = imp.detect_duplicate_blocks(loaded, ['v2', 'v3'], formal, [])
        self.assertEqual(len(rules), 1)
        self.assertEqual(rules[0]['keep']['part'], 'v3')
        self.assertEqual(rules[0]['drop']['part'], 'v2')
        self.assertEqual(rules[0]['rawItems'], 2)


class StyleCollectionTests(unittest.TestCase):
    RULE = {'part': 'N5-regular', 'path': ['画风搜集'], 'hosts': {'画风1': ['负面', '设置'], '画风2': ['负面', '附带tag']}}

    def test_labels_fold_into_their_style_card(self):
        entries = [
            card('画风1', ['画风搜集'], 'artist:a,\nshiny skin,', isNew=True),
            card('负面', ['画风搜集'], 'lowres,'),
            card('设置', ['画风搜集'], 'Steps:  28\nSampler:  k_euler_ancestral (karras) '),
            card('画风2', ['画风搜集'], 'artist:b,'),
            card('负面', ['画风搜集'], 'blurry,'),
            card('附带tag', ['画风搜集'], 'maid,', characterPrompts=[{'label': 'char1', 'prompt': 'girl,'}]),
            card('画风3', ['画风搜集'], 'artist:c,'),
        ]
        folded = imp.fold_style_collection(entries, self.RULE)
        self.assertEqual([e['title'] for e in folded], ['画风1', '画风2', '画风3'])
        self.assertEqual(folded[0]['negative'], 'lowres,')
        self.assertEqual(folded[0]['note'], 'Steps: 28 · Sampler: k_euler_ancestral (karras)')
        self.assertEqual(folded[1]['tags'], 'artist:b,\nmaid,')
        self.assertEqual(folded[1]['characterPrompts'], [{'label': 'char1', 'prompt': 'girl,'}])

    def test_changed_shape_blocks(self):
        entries = [card('画风1', ['画风搜集'], 'artist:a,'), card('负面', ['画风搜集'], 'lowres,')]
        with self.assertRaises(ValueError):
            imp.fold_style_collection(entries, self.RULE)

    def test_detection_skips_view_variant_titles(self):
        parts = {'N5-regular': [card('画风1', ['画风搜集'], 'a,'), card('负面', ['画风搜集'], 'b,')],
                 'N4.5-v4': [card('背面', ['收录'], 'back,'), card('正面', ['收录'], 'front,')]}
        rules, orphans = imp.detect_style_collections(parts)
        self.assertEqual(rules, [{'part': 'N5-regular', 'path': ['画风搜集'], 'hosts': {'画风1': ['负面']}}])
        self.assertEqual(orphans, [])


class AliasAndHalfTests(unittest.TestCase):
    def test_renamed_section_is_detected_and_matches_exactly(self):
        old = [{'id': 'codex_6e699406-0001', 'title': '双人', 'path': ['各种', '2+girl/+1boy系列', '百合'], 'tags': 'a,'},
               {'id': 'codex_6e699406-0002', 'title': '三人', 'path': ['各种', '2+girl/+1boy系列', '百合'], 'tags': 'b,'},
               {'id': 'codex_6e699406-0003', 'title': '留下', 'path': ['各种', '1girl系列'], 'tags': 'c,'}]
        new = [card('双人', ['各种', '2+girl/2+boy系列', '百合'], 'a,'), card('三人', ['各种', '2+girl/2+boy系列', '百合'], 'b,'),
               card('留下', ['各种', '1girl系列'], 'c,')]
        pairs = imp.detect_aliases(old, new)
        self.assertEqual(pairs, [(('各种', '2+girl/+1boy系列'), ('各种', '2+girl/2+boy系列'))])
        result = match.match_entries(imp.aliased(old, pairs), new)
        self.assertEqual(result['summary']['methodCounts'], {'exact_fingerprint': 3})
        self.assertEqual(old[0]['path'][1], '2+girl/+1boy系列')

    def test_halves_follow_old_id_namespaces(self):
        formal = {'suozhang_r18': {'entries': [
            {'id': 'codex_6e699406-0001', 'title': '甲', 'path': ['上'], 'tags': 'a,'},
            {'id': 'codex_8489ac52-0001', 'title': '乙', 'path': ['下'], 'tags': 'b,'}]}}
        parts = {'N4.5-v1': [card('甲', ['上'], 'a,')], 'N4.5-v2': [card('乙', ['下'], 'b,')],
                 'N4.5-v3': [card('全新', ['新'], 'z,')]}
        halves, _ = imp.assign_halves('suozhang_r18', ['N4.5-v1', 'N4.5-v2', 'N4.5-v3'], parts, formal)
        self.assertEqual(halves, {'N4.5-v1': 'upper', 'N4.5-v2': 'lower', 'N4.5-v3': None})


class RemovalTests(unittest.TestCase):
    def rows(self):
        sources = {book: [] for book in imp.BOOKS}
        sources['suozhang_r18'] = [card('留', ['基础'], 'keep,')]
        sources['suozhang_nai5_r18'] = [card('搬走', ['百合'], 'moved,')]
        removed = [
            {'id': 'r-1', 'title': '附录', 'path': ['附录栏目'], 'tags': 'x,'},
            {'id': 'r-2', 'title': '搬走', 'path': ['基础'], 'tags': 'moved,'},
            {'id': 'r-3', 'title': '消失', 'path': ['基础'], 'tags': 'gone,'},
            {'id': 'r-4', 'title': '有图', 'path': ['附录栏目'], 'tags': 'y,', 'image': 'r-4.jpg'},
        ]
        return imp.classify_removals('suozhang_r18', removed, sources, [])

    def test_reasons(self):
        reasons = {row['id']: (row['reason'], row['imaged']) for row in self.rows()}
        self.assertEqual(reasons, {'r-1': ('sectionGone', False), 'r-2': ('movedTo', False),
                                   'r-3': ('unexplained', False), 'r-4': ('sectionGone', True)})

    def test_only_approved_removals_pass(self):
        rows = self.rows()
        config = {'removals': {'suozhang_r18': {'retiredSections': ['附录栏目'],
                                                'movedTo': {'suozhang_nai5_r18': ['r-2']}}}}
        with self.assertRaises(ValueError):  # pictured entry needs explicit approval
            imp.check_removals(config, 'suozhang_r18', rows)
        config['removals']['suozhang_r18']['approvedImagedIds'] = ['r-4']
        with self.assertRaises(ValueError):  # unexplained removal
            imp.check_removals(config, 'suozhang_r18', rows)
        config['removals']['suozhang_r18']['approvedIds'] = ['r-3']
        imp.check_removals(config, 'suozhang_r18', rows)


class ManualMatchAndDriftTests(unittest.TestCase):
    def test_manual_match_pins_rewritten_card_to_old_id(self):
        old = [{'id': 'b-0070', 'title': '旧无标题', 'path': ['收录'], 'tags': 'a,b,c,'}]
        new = [card('新标题', ['收录'], 'a,b,c,d,')]
        config = {'manualMatches': {'demo': [{'oldId': 'b-0070',
                                              'newSignatureSha256': match.override_signature_sha256(new[0])}]}}
        forced = imp.resolve_manual_matches(config, 'demo', old, new)
        result = match.match_entries(old, new, forced_pairs=forced)
        self.assertEqual(result['matches'][0]['old']['id'], 'b-0070')
        self.assertEqual(result['summary']['methodCounts'], {'manual_override': 1})
        with self.assertRaises(ValueError):  # source text changed since the decision
            imp.resolve_manual_matches(config, 'demo', old, [card('新标题', ['收录'], 'changed,')])

    def test_only_listed_local_edits_pass_replay(self):
        baseline = {'suozhang_r18': {'summary': {'strictReplayPass': False}, 'unmatchedOld': [], 'unmatchedNew': 0,
                                     'drift': [{'id': 'x-1', 'fields': ['tags']}]}}
        self.assertEqual(imp.replay_failures({}, baseline), ['suozhang_r18'])
        config = {'acceptedDrift': {'suozhang_r18': [{'id': 'x-1', 'fields': ['tags']}]}}
        self.assertEqual(imp.replay_failures(config, baseline), [])
        stale = {'acceptedDrift': {'suozhang_r18': [{'id': 'x-1', 'fields': ['tags']}, {'id': 'x-2', 'fields': ['title']}]}}
        self.assertEqual(imp.replay_failures(stale, baseline), ['suozhang_r18'])


class SourceSplitTests(unittest.TestCase):
    def rule(self):
        return {'part': 'N5-regular', 'lineSha256': imp.sha256_text('long untitled paragraph,'),
                'title': '收录（无标题1）', 'isNew': False, 'note': 'demo'}

    def test_registered_paragraph_is_split_back(self):
        entries = [card('新卡', ['收录'], 'new card,\nlong untitled paragraph,', isNew=True),
                   card('旧卡', ['收录'], 'old,')]
        result = imp.apply_source_splits('N5-regular', entries, [self.rule()])
        self.assertEqual([(e['title'], e['tags'], e['isNew']) for e in result],
                         [('新卡', 'new card,', True), ('收录（无标题1）', 'long untitled paragraph,', False), ('旧卡', 'old,', False)])

    def test_record_stops_when_author_fixed_it(self):
        standalone = [card('有了标题', ['收录'], 'long untitled paragraph,')]
        with self.assertRaises(ValueError):
            imp.apply_source_splits('N5-regular', standalone, [self.rule()])
        with self.assertRaises(ValueError):
            imp.apply_source_splits('N5-regular', [card('新卡', ['收录'], 'edited paragraph,')], [self.rule()])
        self.assertEqual(imp.apply_source_splits('N4.5-regular', standalone, [self.rule()]), standalone)


class SpecialChapterTests(unittest.TestCase):
    RULE = {'part': 'N4.5-regular', 'chapter': '自然语言', 'titlePrefix': '自然语言'}

    def document(self, lines):
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        doc = Document()

        def heading(text):
            paragraph = doc.add_paragraph(text)
            outline = OxmlElement('w:outlineLvl')
            outline.set(qn('w:val'), '0')
            paragraph._p.get_or_add_pPr().append(outline)

        heading('场景')
        doc.add_paragraph('普通卡')
        heading('自然语言')
        for line in lines:
            doc.add_paragraph(line)
        heading('群友处收录')
        return doc

    def test_template_blocks_become_numbered_cards(self):
        doc = self.document(['原Tag：', 'a calm lake at dawn,', '翻译：', '黎明时平静的湖面', '',
                             '原Tag：', 'a city street,', 'neon lights,', '翻译：', '城市街道', 'Tag版本改编后：', 'city, neon,',
                             '原Tag：', '翻译：', '原Tag：', '翻译：'])
        items, later, stats = imp.parse_template_chapter(doc, self.RULE)
        self.assertEqual([(i['title'], i['tags'], i.get('note')) for i in items], [
            ('自然语言（1）', 'a calm lake at dawn,', '黎明时平静的湖面'),
            ('自然语言（2）', 'a city street,\nneon lights,', '城市街道'),
            ('自然语言（2）：改编后', 'city, neon,', None)])
        self.assertEqual(later, ['群友处收录'])
        self.assertEqual((stats['filled'], stats['emptyTemplates'], stats['adapted']), (2, 2, 1))
        self.assertFalse(any(i['isNew'] for i in items))

    def test_translation_without_original_stops(self):
        with self.assertRaises(ValueError):
            imp.parse_template_chapter(self.document(['原Tag：', '翻译：', '只有翻译']), self.RULE)

    def test_items_go_back_to_the_chapter_position(self):
        raw = [card('甲', ['场景'], 'a,'), card('乙', ['群友处收录'], 'b,')]
        merged = imp.insert_chapter_items(raw, [card('自然语言（1）', ['自然语言'], 'x,')], ['群友处收录', '页尾'])
        self.assertEqual([e['title'] for e in merged], ['甲', '自然语言（1）', '乙'])

    def test_site_titles_survive_and_new_cards_continue_numbering(self):
        config = {'specialChapters': [self.RULE]}
        formal = {book: {'entries': []} for book in imp.BOOKS}
        formal['suozhang']['entries'] = [
            {'id': 's-1', 'title': '晨雾湖面', 'path': ['自然语言'], 'tags': 'a calm lake at dawn,'},
            {'id': 's-2', 'title': '自然语言（2）', 'path': ['自然语言'], 'tags': 'a city street,'}]
        sources = {book: [] for book in imp.BOOKS}
        sources['suozhang'] = [card('自然语言（1）', ['自然语言'], 'a calm lake at dawn,'),
                               card('自然语言（2）', ['自然语言'], 'a new forest,'),
                               card('自然语言（2）：改编后', ['自然语言'], 'forest,'),
                               card('自然语言（3）', ['自然语言'], 'a city street,')]
        imp.own_chapter_titles(config, formal, sources)
        self.assertEqual([e['title'] for e in sources['suozhang']],
                         ['晨雾湖面', '自然语言（3）', '自然语言（3）：改编后', '自然语言（2）'])

    def test_replay_ignores_site_owned_titles(self):
        config = {'specialChapters': [self.RULE]}
        old = [{'id': 's-1', 'title': '起好的名字', 'path': ['自然语言'], 'tags': 'a,'}]
        snapshot = [card('自然语言（1）', ['自然语言'], 'a,')]
        owned = imp.chapter_paths(config, 'suozhang')
        blank = [dict(e, title='') if tuple(e['path']) in owned else e for e in old]
        result = match.match_entries(blank, [dict(snapshot[0], title='')])
        self.assertTrue(result['summary']['strictReplayPass'])


class ReportHintTests(unittest.TestCase):
    def test_lines_after_image_text_are_flagged(self):
        from docx import Document
        doc = Document()
        lines = ['吃下后迷幻', '1girl,drugged,', '迷幻时侵犯', 'char1：girl,',
                 '小穴黏在灯柱上', '1girl,solo,text: 唔……是不小心的！', '帮帮我!', 'char1：girl,solo,',
                 '招募', 'char1：parchment,text:', '【募集】', '现招募志愿者。', '男女不限,有无经验均可。', '※请服从指示。', 'char2：girl,']
        for line in lines:
            doc.add_paragraph(line)
        paragraphs, at = doc.paragraphs, {line: i for i, line in enumerate(lines)}
        self.assertFalse(imp.after_image_text(paragraphs, at['迷幻时侵犯'], set()))
        self.assertTrue(imp.after_image_text(paragraphs, at['帮帮我!'], set()))
        run = {at['【募集】'], at['现招募志愿者。'], at['※请服从指示。']}
        self.assertTrue(imp.after_image_text(paragraphs, at['现招募志愿者。'], run))
        self.assertTrue(imp.after_image_text(paragraphs, at['※请服从指示。'], run))
        self.assertFalse(imp.after_image_text(paragraphs, at['※请服从指示。'], set()))  # 「现招募」 alone reads as a title

    def test_copies_and_followed_moves(self):
        old = [{'id': 'a-1', 'title': '组', 'path': ['画师'], 'tags': 'artist:a,', 'image': 'a-1.jpg'},
               {'id': 'a-2', 'title': '留下', 'path': ['旧目录'], 'tags': 'keep,'}]
        new = [card('组', ['画师'], 'artist:a,'), card('组', ['画师'], 'artist:a,'), card('留下', ['新目录'], 'keep,')]
        result = match.match_entries(old, new)
        self.assertEqual(imp.copied_additions(new, result),
                         [{'index': 1, 'title': '组', 'path': ['画师'], 'oldId': 'a-1', 'oldWentTo': 0, 'oldPath': ['画师'], 'imaged': True}])
        self.assertEqual(imp.followed_moves(result), [{'from': ['旧目录'], 'to': ['新目录'], 'count': 1}])
        copy = imp.copied_additions(new, result)[0]
        self.assertIn('同一目录里出现了两份', imp.describe_copy('demo', copy))
        self.assertIn('别的目录', imp.describe_copy('demo', dict(copy, path=['自然语言'])))

    def test_prose_outcomes_name_the_cards(self):
        path = ['涩涩', '杂项']
        rows = [{'part': 'N5-upper', 'paragraph': p, 'path': path, 'text': t, 'afterImageText': h}
                for p, t, h in ((10, '帮帮我!', True), (20, '续一句', True), (30, '空标题', False), (40, '丢了', False))]
        entries = [{'id': 'b-1', 'title': '灯柱', 'path': path, 'tags': 'text: 唔,',
                    'characterPrompts': [{'prompt': 'girl,\n续一句'}]},
                   {'id': 'b-2', 'title': '帮帮我!', 'path': path, 'tags': ''}]
        built = {'loaded': {'N5-upper': {'doc': None}}, 'applied': {'suozhang_nai5_r18': {'entries': entries}}}
        config = {'proseParagraphs': {'N5-upper': {'20': 'x', '40': 'y'}}}
        with mock.patch.object(imp, 'prose_candidates', return_value=rows):
            outcomes = imp.prose_outcomes(config, built)
        self.assertEqual([(o['outcome'], o.get('card')) for o in outcomes],
                         [('ownCard', ['b-2', '帮帮我!']), ('merged', ['b-1', '灯柱']), ('noCard', None), ('missing', None)])
        self.assertEqual(outcomes[0]['cardAbove'], ['b-1', '灯柱'])
        self.assertIn('自成一卡', imp.render_prose_outcomes(outcomes))

    def test_audit_flags_catch_swallowed_titles(self):
        swallowed = card('逼近注视', ['搞怪'], 'white background,', characterPrompts=[
            {'label': 'char1', 'prompt': 'girl,'}, {'label': 'char2', 'prompt': 'girl,\n这里跟你无关，继续往下滑\nhorror,'},
            {'label': 'char1', 'prompt': 'girl,'}])
        self.assertEqual(imp.audit_flags(swallowed), ['角色框编号重来', '正文里夹着中文短句'])
        self.assertEqual(imp.audit_flags(card('搞怪（无标题1）', ['搞怪'], '烧死他，用文火')), ['无标题', '正文里夹着中文短句'])
        self.assertEqual(imp.audit_flags(card('正常', ['搞怪'], '1girl,smile,', characterPrompts=[
            {'label': 'char1', 'prompt': 'girl,'}, {'label': 'char2', 'prompt': 'boy,'}])), [])
        old = {'id': 'b-1', 'title': '旧', 'path': ['搞怪'], 'tags': 'a,'}
        built = {'results': {book: {'match': match.match_entries([], [])} for book in imp.BOOKS},
                 'applied': {book: {'entries': []} for book in imp.BOOKS}}
        built['results']['suozhang']['match'] = match.match_entries([old], [card('旧', ['搞怪'], 'a,b,')])
        built['applied']['suozhang']['entries'] = [dict(old, tags='a,b,'), dict(swallowed, id='b-2')]
        report = imp.changes_report(built)['suozhang']
        self.assertEqual([r['id'] for r in report['added']], ['b-2'])
        self.assertEqual([(r['id'], r['fields']) for r in report['changed']], [('b-1', ['tags'])])
        self.assertIn('⚑ 角色框编号重来', imp.render_changes(imp.changes_report(built)))

    def test_content_signature_ignores_ids(self):
        entry = card('甲', ['场景'], 'a,')
        self.assertEqual(imp.content_signature(dict(entry, id='x-1')), imp.content_signature(dict(entry, id='x-2')))
        self.assertNotEqual(imp.entry_signature(dict(entry, id='x-1')), imp.entry_signature(dict(entry, id='x-2')))


class WordDocMixin:
    """Writes a minimal codex Word file: a TOC, a heading 「场景」, then ``lines``.

    A line starting with ``#`` becomes a level-2 outline heading.
    """

    def write_doc(self, folder, lines=None):
        from docx import Document
        from docx.enum.style import WD_STYLE_TYPE
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        doc = Document()
        doc.add_paragraph('目录', style=doc.styles.add_style('toc 1', WD_STYLE_TYPE.PARAGRAPH))

        def heading(text, level):
            paragraph = doc.add_paragraph(text)
            outline = OxmlElement('w:outlineLvl')
            outline.set(qn('w:val'), str(level))
            paragraph._p.get_or_add_pPr().append(outline)

        heading('场景', 0)
        doc.add_paragraph('')
        for line in lines or self.LINES:
            if line.startswith('#'):
                heading(line[1:], 1)
            else:
                doc.add_paragraph(line)
        path = Path(folder) / 'n5.docx'
        doc.save(path)
        return path, {line: i for i, line in enumerate(p.text for p in Document(path).paragraphs) if line}


class TitleParagraphTests(WordDocMixin, unittest.TestCase):
    LINES = ['烧死他，用文火', 'char1：girl,', '', '逼近注视', 'white background,', 'char1：girl,', '',
             '这里跟你无关，继续往下滑', 'horror (theme),', 'char1：girl,']

    def test_registered_lines_become_titles(self):
        with tempfile.TemporaryDirectory() as folder:
            path, at = self.write_doc(folder)
            plain = imp.load_part(path, 'N5-regular')['raw']
            self.assertEqual([e['title'] for e in plain], ['(未命名)', '逼近注视'])
            titles = {str(at[t]): imp.sha256_text(t) for t in ('烧死他，用文火', '这里跟你无关，继续往下滑')}
            fixed = imp.load_part(path, 'N5-regular', titles=titles)['raw']
            self.assertEqual([e['title'] for e in fixed], ['烧死他，用文火', '逼近注视', '这里跟你无关，继续往下滑'])
            with self.assertRaises(ValueError):  # the paragraph moved or changed
                imp.load_part(path, 'N5-regular', titles={str(at['逼近注视']): imp.sha256_text('烧死他，用文火')})
            with self.assertRaises(ValueError):  # one paragraph cannot be both
                imp.load_part(path, 'N5-regular', prose=titles, titles=titles)

    def test_detection_and_relocation(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as folder:
            path, at = self.write_doc(folder)
            docs = {'N5-regular': Document(path)}
            rows = imp.title_candidates(docs, {})
            self.assertEqual([r['text'] for r in rows], ['烧死他，用文火', '这里跟你无关，继续往下滑'])
            previous = {'titleParagraphs': {'N5-regular': {'99': imp.sha256_text('这里跟你无关，继续往下滑')}}}
            found, missing = imp.relocate_titles(previous, docs)
            self.assertEqual((found, missing), ({'N5-regular': {str(at['这里跟你无关，继续往下滑']):
                                                                 imp.sha256_text('这里跟你无关，继续往下滑')}}, []))
            self.assertEqual([r['text'] for r in imp.title_candidates(docs, found)], ['烧死他，用文火'])

    def test_title_broken_inside_brackets_is_detected(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as folder:
            path, _ = self.write_doc(folder, ['前后夹击', 'solo,', 'char1：girl,', '前后夹击（口/穴，桌底', '）1koma,', 'char1：girl,'])
            rows = imp.title_candidates({'N5-upper': Document(path)}, {})
            self.assertEqual([r['text'] for r in rows], ['前后夹击（口/穴，桌底'])


class TranslationParagraphTests(WordDocMixin, unittest.TestCase):
    PAIRS = ['仰头吞剑', '1girl,sword,', '',
             '1boy,femboy,slave market,', '一个男孩，伪娘，站在昏暗的地下奴隶市场里，戴着白色头巾，内里是蓝色的头发，红色的眼睛，明亮的瞳孔。', '',
             'a wolf-boy trapped in a wall,', 'dark room,', '一个狼孩被困在墙壁的缝隙中，只能透过木洞看到臀部，场景设定在一个昏暗的地下奴隶市场。', '黑暗的房间。', '',
             '连体泳装侵犯', 'swimsuit,']

    def test_translations_split_cards_and_become_notes(self):
        with tempfile.TemporaryDirectory() as folder:
            path, at = self.write_doc(folder, self.PAIRS)
            plain = imp.load_part(path, 'N4.5-v4')['raw']
            self.assertEqual(plain[0]['title'], '仰头吞剑')
            self.assertIn('1boy,femboy', plain[0]['tags'])  # the bug: swallowed by the card above
            lines = [t for t in self.PAIRS if t.startswith(('一个', '黑暗'))]
            registered = {str(at[t]): imp.sha256_text(t) for t in lines}
            fixed = imp.load_part(path, 'N4.5-v4', translations=registered)['raw']
            self.assertEqual([(e['title'], e['tags'], e.get('note', '')[:4]) for e in fixed], [
                ('仰头吞剑', '1girl,sword,', ''),
                ('(未命名)', '1boy,femboy,slave market,', '一个男孩'),
                ('(未命名)', 'a wolf-boy trapped in a wall,\ndark room,', '一个狼孩'),
                ('连体泳装侵犯', 'swimsuit,', '')])
            self.assertTrue(fixed[2]['note'].endswith('\n黑暗的房间。'))
            with self.assertRaises(ValueError):  # a translation needs prompt lines right above it
                imp.load_part(path, 'N4.5-v4', translations={str(at['仰头吞剑']): imp.sha256_text('仰头吞剑')})

    def test_translation_under_a_title_stays_with_that_card(self):
        with tempfile.TemporaryDirectory() as folder:
            path, at = self.write_doc(folder, ['脑交', 'venus,kneeling,', '维纳斯跪着，头部向上倾斜，角色的姿势是顺从的，头部向后倾斜，嘴巴微微张开，跪在地上。'])
            line = '维纳斯跪着，头部向上倾斜，角色的姿势是顺从的，头部向后倾斜，嘴巴微微张开，跪在地上。'
            raw = imp.load_part(path, 'N4.5-v4', translations={str(at[line]): imp.sha256_text(line)})['raw']
            self.assertEqual([(e['title'], e['tags'], bool(e.get('note'))) for e in raw], [('脑交', 'venus,kneeling,', True)])

    ONE = '一个男孩站在昏暗的地下奴隶市场里，戴着白色头巾，内里是蓝色的头发，红色的眼睛，明亮的瞳孔，很害怕。'
    TWO = '一个狼孩被困在墙壁的缝隙中，只能透过木洞看到臀部，场景设定在一个昏暗的地下奴隶市场，很屈辱。'

    def registered(self, at, *lines):
        return {str(at[line]): imp.sha256_text(line) for line in lines}

    def test_pairs_without_blank_line_between_stay_apart(self):
        with tempfile.TemporaryDirectory() as folder:
            path, at = self.write_doc(folder, ['卡甲', 'a,', '', 'eng one,', self.ONE, 'eng two,', self.TWO])
            raw = imp.load_part(path, 'N4.5-v4', translations=self.registered(at, self.ONE, self.TWO))['raw']
            self.assertEqual([(e['title'], e['tags']) for e in raw],
                             [('卡甲', 'a,'), ('(未命名)', 'eng one,'), ('(未命名)', 'eng two,')])

    def test_prompt_search_stops_at_a_heading(self):
        with tempfile.TemporaryDirectory() as folder:  # a heading such as 「1girl,1boy系列」 reads as prompt text
            path, at = self.write_doc(folder, ['卡甲', 'a,', '', 'b,', '#1girl,1boy系列', 'eng one,', self.ONE])
            raw = imp.load_part(path, 'N4.5-v4', translations=self.registered(at, self.ONE))['raw']
            self.assertEqual([(e['title'], e['tags'], e['path'][-1]) for e in raw],
                             [('卡甲', 'a,\nb,', '场景'), ('(未命名)', 'eng one,', '1girl,1boy系列')])

    def test_prompt_under_a_dropped_note_line_starts_a_card(self):
        note = 'This is a long english explanation line without any tag syntax at all'
        with tempfile.TemporaryDirectory() as folder:
            path, at = self.write_doc(folder, ['卡甲', 'a,', note, 'eng one,', self.ONE])
            raw = imp.load_part(path, 'N4.5-v4', translations=self.registered(at, self.ONE))['raw']
            self.assertEqual([(e['title'], e['tags']) for e in raw], [('卡甲', 'a,'), ('(未命名)', 'eng one,')])

    def test_survey_drops_only_the_translation_blocks_that_no_longer_fit(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as folder:
            path, at = self.write_doc(folder, ['卡甲', 'a,', '', 'eng one,', self.ONE, '', '标题乙', self.TWO])
            translations = {'N4.5-v4': self.registered(at, self.ONE, self.TWO)}
            dropped = imp.prune_translations({'N4.5-v4': Document(path)}, {}, {}, translations)
            self.assertEqual(dropped, {'N4.5-v4': [at[self.TWO]]})
            self.assertEqual(translations, {'N4.5-v4': self.registered(at, self.ONE)})

    def test_relocated_by_hash(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as folder:
            path, at = self.write_doc(folder, self.PAIRS)
            previous = {'translationParagraphs': {'N4.5-v4': {'7': imp.sha256_text(self.PAIRS[4]), '9': 'gone'}}}
            found, missing = imp.relocate_translations(previous, {'N4.5-v4': Document(path)})
            self.assertEqual((found, missing), ({'N4.5-v4': {str(at[self.PAIRS[4]]): imp.sha256_text(self.PAIRS[4])}}, ['gone']))

    def test_source_notes_follow_the_source(self):
        built = [dict(card('甲', ['收录'], 'a,'), id='x-1', note='网站加的'),
                 dict(card('收录（无标题1）', ['收录'], 'b,'), id='x-2', note='旧翻译'),
                 dict(card('自然语言（1）', ['自然语言'], 'c,'), id='x-3', note='作者删了')]
        sources = [card('甲', ['收录'], 'a,'), card('收录（无标题1）', ['收录'], 'b,', note='新翻译'),
                   card('自然语言（1）', ['自然语言'], 'c,')]
        noted = imp.refresh_source_notes('demo', built, sources, {('自然语言',)})
        self.assertEqual(noted, {'x-2'})
        self.assertEqual([e.get('note') for e in built], ['网站加的', '新翻译', None])
        with self.assertRaises(ValueError):  # built entries must line up with the source
            imp.refresh_source_notes('demo', built[:2], sources, set())

    def test_note_changes_are_listed(self):
        old = {'id': 'x-1', 'title': '收录（无标题1）', 'path': ['收录'], 'tags': 'b,'}
        built = {'results': {book: {'match': match.match_entries([], [])} for book in imp.BOOKS},
                 'applied': {book: {'entries': []} for book in imp.BOOKS},
                 'formal': {book: {'entries': []} for book in imp.BOOKS}}
        built['results']['suozhang']['match'] = match.match_entries([old], [card('收录（无标题1）', ['收录'], 'b,')])
        built['applied']['suozhang']['entries'] = [dict(old, note=self.ONE)]
        built['formal']['suozhang']['entries'] = [old]
        self.assertEqual([(r['id'], r['fields']) for r in imp.changes_report(built)['suozhang']['changed']],
                         [('x-1', ['note'])])

    def test_audit_flags_for_sentence_titles_and_translated_cards(self):
        self.assertEqual(imp.audit_flags(card(self.ONE, ['收录'], 'a,')), ['标题是长句'])
        self.assertEqual(imp.audit_flags(card('收录（无标题1）', ['收录'], 'eng one,', note=self.ONE)), [])
        self.assertEqual(imp.audit_flags(card('收录（无标题1）', ['收录'], 'eng one,')), ['无标题'])

    def test_long_sentences_read_as_titles_are_reported(self):
        from docx import Document
        self.assertTrue(imp.sentence_title('一个男孩，伪娘，戴着白色头巾，内里是蓝色的头发，红色的眼睛，明亮的瞳孔。他像商品一样站着。'))
        self.assertFalse(imp.sentence_title('扼喉背后位【推荐画风：[[artist:a]],[[artist:b]],[[artist:c]],[[artist:d]],[[artist:e]]】'))
        with tempfile.TemporaryDirectory() as folder:
            path, at = self.write_doc(folder, self.PAIRS)
            rows = imp.long_title_candidates({'N4.5-v4': Document(path)}, set())
            self.assertEqual([(r['paragraph'], r['afterPrompt']) for r in rows],
                             [(at[self.PAIRS[4]], True), (at[self.PAIRS[8]], True), (at[self.PAIRS[9]], False)])


class JoinedTitleTests(WordDocMixin, unittest.TestCase):
    LINES = ['半拉开外套版本', 'eyes focus,shy,', '附带两个动作', '饮水：towel around neck,holding bottle,',
             '', '小巷撩起衣服诱惑小正太1girl,solo,sunlight,amusement\npark,best illustration,', 'char1：girl,smile,']

    def registered(self, at, line, cut):
        return {str(at[line]): {'sha256': imp.sha256_text(line), 'at': cut}}

    def test_title_written_into_the_prompt_line_is_cut_off(self):
        with tempfile.TemporaryDirectory() as folder:
            path, at = self.write_doc(folder)
            plain = imp.load_part(path, 'N4.5-v4')['raw']
            self.assertNotIn('饮水', [e['title'] for e in plain])
            joined = {**self.registered(at, self.LINES[3], 2), **self.registered(at, self.LINES[5], 11)}
            raw = imp.load_part(path, 'N4.5-v4', joined=joined)['raw']
            self.assertEqual([(e['title'], e['tags']) for e in raw], [
                ('半拉开外套版本', 'eyes focus,shy,'),
                ('饮水', 'towel around neck,holding bottle,'),
                ('小巷撩起衣服诱惑小正太', '1girl,solo,sunlight,amusement\npark,best illustration,\nchar1：girl,smile,')])

    def test_cut_must_leave_a_title_and_a_prompt(self):
        with tempfile.TemporaryDirectory() as folder:
            path, at = self.write_doc(folder)
            with self.assertRaises(ValueError):  # cut in the middle of the prompt
                imp.load_part(path, 'N4.5-v4', joined=self.registered(at, self.LINES[5], 20))
            with self.assertRaises(ValueError):  # the paragraph changed
                imp.load_part(path, 'N4.5-v4', joined={str(at[self.LINES[3]]): {'sha256': 'x', 'at': 2}})

    def test_n45_registrations_are_found_again(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as folder:
            path, at = self.write_doc(folder, ['抱枕双面', '角色tag+,arm up,armpits,ass,back,bare legs,barefoot,bed sheet,breasts', *self.LINES])
            docs = {'N4.5-regular': Document(path)}
            prose_line = '角色tag+,arm up,armpits,ass,back,bare legs,barefoot,bed sheet,breasts'
            found, _ = imp.relocate_prose({'proseParagraphs': {'N4.5-regular': {'1': imp.sha256_text(prose_line)}}}, docs)
            self.assertEqual(found, {'N4.5-regular': {str(at[prose_line]): imp.sha256_text(prose_line)}})
            previous = {'joinedTitles': {'N4.5-regular': {'3': {'sha256': imp.sha256_text(self.LINES[3]), 'at': 2}}}}
            found, missing = imp.relocate_joined_titles(previous, docs)
            self.assertEqual((found, missing), ({'N4.5-regular': self.registered(at, self.LINES[3], 2)}, []))


class DroppedLineTests(WordDocMixin, unittest.TestCase):
    PROSE = 'shows a witch standing before a glowing magic circle in her dark study'
    LINES = ['女巫触手召唤败北', '3koma,multiple views,', PROSE, '', '穿鞋袜出门漫画', '{four-panel comics},',
             'first panel: wear white pantyhose', 'second panel: wear shoes', 'fourth panel: going out,location,']

    def test_dropped_english_and_section_titles_are_reported(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as folder:
            path, at = self.write_doc(folder)
            docs = {'N4.5-v3': Document(path)}
            self.assertEqual([r['paragraph'] for r in imp.dropped_lines(docs, set())], [at[self.PROSE]])
            sections = [r['text'] for r in imp.long_title_candidates(docs, set()) if r['section']]
            self.assertEqual(sections, ['first panel: wear white pantyhose', 'second panel: wear shoes'])
            raw = imp.load_part(path, 'N4.5-v3', prose={str(at[self.PROSE]): imp.sha256_text(self.PROSE)})['raw']
            self.assertEqual(raw[0]['tags'], '3koma,multiple views,\n' + self.PROSE)

    def test_dropped_prose_is_found_again_only_in_the_same_model(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as folder:
            path, at = self.write_doc(folder)
            doc = Document(path)
            registered = {'proseParagraphs': {'N4.5-v3': {'1': imp.sha256_text(self.PROSE)}}}
            found, missing = imp.relocate_prose(registered, {'N4.5-v4': doc})
            self.assertEqual((found, missing), ({'N4.5-v4': {str(at[self.PROSE]): imp.sha256_text(self.PROSE)}}, []))
            found, missing = imp.relocate_prose(registered, {'N5-upper': doc})  # an N4.5 line is never moved into N5
            self.assertEqual((found, missing), ({}, [imp.sha256_text(self.PROSE)]))

    def test_cut_inside_an_english_word_is_refused(self):
        line = '小巷诱惑小正太1girl,solo,sunlight,'
        with tempfile.TemporaryDirectory() as folder:
            path, at = self.write_doc(folder, ['卡甲', 'a,', '', line])
            for cut in (8, 9):  # 「…小正太1」 / 「…小正太1g」
                with self.assertRaises(ValueError):
                    imp.load_part(path, 'N4.5-v4', joined={str(at[line]): {'sha256': imp.sha256_text(line), 'at': cut}})
            raw = imp.load_part(path, 'N4.5-v4', joined={str(at[line]): {'sha256': imp.sha256_text(line), 'at': 7}})['raw']
            self.assertEqual(raw[-1]['title'], '小巷诱惑小正太')


class RehearsalGuardTests(unittest.TestCase):
    def test_survey_refuses_to_write_into_an_applied_archive(self):
        with tempfile.TemporaryDirectory() as source, tempfile.TemporaryDirectory() as out:
            (Path(source) / 'N5所长常规NovelAI个人法典（2026.10.13版）.docx').write_bytes(b'')
            (Path(out) / 'applied.json').write_text('{}', encoding='utf-8')
            with self.assertRaises(ValueError):
                imp.survey(source, out=Path(out))

    def test_rehearsal_start_config_loads(self):
        config = imp.load_config(imp.VERSIONS_DIR / 'rehearsal' / '1013-previous.json')
        self.assertEqual(config['version'], '2026.9.25')
        self.assertEqual([r['chapter'] for r in config['specialChapters']], ['自然语言'])
        self.assertNotIn('sourceSplits', config)


class ConfigTests(unittest.TestCase):
    def write(self, value):
        handle = tempfile.NamedTemporaryFile('w', suffix='.json', delete=False, encoding='utf-8')
        json.dump(value, handle, ensure_ascii=False)
        handle.close()
        return Path(handle.name)

    def test_half_rules(self):
        base = {'version': '2026.1.1', 'historical': True, 'outDir': 'output/x'}
        with self.assertRaises(ValueError):
            imp.load_config(self.write({**base, 'parts': {'N4.5-v1': {'half': None}}}))
        with self.assertRaises(ValueError):
            imp.load_config(self.write({**base, 'parts': {'N5-regular': {'half': 'upper'}}}))
        loaded = imp.load_config(self.write({**base, 'parts': {'N4.5-v1': {'half': 'upper'}}}))
        self.assertEqual(loaded['parts']['N4.5-v1']['half'], 'upper')

    def test_shipped_configs_load(self):
        for path in sorted(imp.VERSIONS_DIR.glob('*.json')):
            config = imp.load_config(path)
            if not config.get('historical'):
                self.assertEqual(set(config['expected']), set(imp.BOOKS))

    def test_overrides_install_without_conflict(self):
        config = {'version': '2099.1.1', 'newOverrides': {'demo': {'demo-0001': {'signatureSha256': 'ab', 'isNew': True}}}}
        imp.install_overrides(config)
        imp.install_overrides(config)
        with self.assertRaises(ValueError):
            imp.install_overrides({**config, 'newOverrides': {'demo': {}}})
        match.AUDITED_NEW_OVERRIDES.pop(('demo', '2099.1.1'))


if __name__ == '__main__':
    unittest.main()
